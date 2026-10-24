"""Lập "Bảng kê 01 - Giảm thuế GTGT" (mẫu HTKK theo NQ142/2024/QH15, TT80) từ hóa đơn đã kéo về.

Điền đúng file mẫu .xls của HTKK (sheet "Bang ke"):
- Phần I : hàng hóa, dịch vụ MUA VÀO có thuế suất 8% (cột tên hàng, giá trị chưa thuế, thuế GTGT).
- Phần II: hàng hóa, dịch vụ BÁN RA có thuế suất 8% (giá trị chưa thuế, thuế suất theo quy định 10,
  thuế suất sau giảm 8, thuế GTGT được giảm = giá trị x 2%).
Mỗi kỳ kê khai (quý hoặc tháng) một file; các dòng cùng tên hàng được cộng gộp.
Kèm file BangKe01_ChiTiet.xlsx liệt kê từng hóa đơn, từng dòng hàng đã cộng vào bảng kê để đối chiếu.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from .misa import TEMPLATES_DIR, _collect, _load_raw_index, _norm_name, skip_reason
from .thuesuat import breakdown, bucket_of

log = logging.getLogger("hddt")

TEMPLATE = TEMPLATES_DIR / "bang_ke_01_giam_thue_gtgt.xls"
FOLDER = "bang_ke_giam_thue"
DETAIL_FILE = "BangKe01_ChiTiet.xlsx"
SHEET = "Bang ke"
RATE = "8%"
TS_QUY_DINH = 10  # thuế suất theo quy định
TS_SAU_GIAM = 8   # = 10 x 80%
MAX_TEN = 150


@dataclass
class Item:
    ten: str
    gia_tri: float = 0.0
    thue: float = 0.0


@dataclass
class BangKeResult:
    folder: str = ""
    files: list[str] = field(default_factory=list)
    detail_path: str = ""
    periods: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def summary(self) -> str:
        if not self.files:
            return "không có hóa đơn thuế suất 8%"
        return ", ".join(p["ky"] for p in self.periods)


def period_of(d: datetime | None, ky: str) -> tuple[tuple[int, int], str, str] | None:
    """Trả về (khóa sắp xếp, nhãn kỳ, phần tên file) của ngày lập hóa đơn."""
    if d is None:
        return None
    if ky == "thang":
        return (d.year, d.month), f"Tháng {d.month:02d}/{d.year}", f"Thang{d.month:02d}_{d.year}"
    q = (d.month - 1) // 3 + 1
    return (d.year, q), f"Quý {q}/{d.year}", f"Quy{q}_{d.year}"


def _cut(text: str, n: int = MAX_TEN) -> str:
    text = re.sub(r"\s+", " ", (text or "").strip())
    return text if len(text) <= n else text[: n - 1] + "…"


# ----------------------------------------------------------------- thu thập số liệu


def collect(out_dir: Path, ky: str = "quy"):
    """Trả về (periods, detail, skipped).

    periods: {khóa kỳ: {"label", "file", "purchase": {tên chuẩn hóa: Item}, "sold": {...}}}
    """
    periods: dict[tuple[int, int], dict[str, Any]] = {}
    detail: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    for direction in ("purchase", "sold"):
        raw_index = _load_raw_index(out_dir, direction)
        chieu = "Mua vào" if direction == "purchase" else "Bán ra"
        for stem, path, inv, _has_xml in _collect(out_dir, raw_index, direction):
            raw = raw_index.get(stem, {})
            so = str(raw.get("shdon") or inv.shdon or "")
            kh = inv.khhdon or str(raw.get("khhdon") or "")
            doi_tac = (inv.nb_ten if direction == "purchase" else (inv.nm_ten or inv.nm_hvtn)) or ""
            mst = (inv.nb_mst if direction == "purchase" else inv.nm_mst) or ""
            label = f"{inv.khmshdon}{kh} số {so}"
            sums, _src = breakdown(raw, inv)
            lines = [ln for ln in inv.lines if ln.tchat != "4" and bucket_of(ln.thue_suat) == RATE]
            if RATE not in sums and not lines:
                continue
            reason = skip_reason(raw)
            if reason:
                skipped.append({"chieu": chieu, "hd": label, "ngay": inv.ngay_lap, "doi_tac": doi_tac, "ly_do": reason})
                continue
            p = period_of(inv.ngay_lap, ky)
            if p is None:
                skipped.append({"chieu": chieu, "hd": label, "ngay": None, "doi_tac": doi_tac, "ly_do": "Không đọc được ngày lập"})
                continue
            key, plabel, pfile = p
            per = periods.setdefault(key, {"label": plabel, "file": pfile, "purchase": {}, "sold": {}})
            bucket = per[direction]
            base = {"ky": plabel, "chieu": chieu, "hd": label, "ngay": inv.ngay_lap, "mst": mst, "doi_tac": doi_tac}

            entries: list[tuple[str, float, float, str]] = []
            for ln in lines:
                sign = -1 if ln.tchat == "3" else 1
                gt = sign * (ln.thanh_tien or 0)
                thue = sign * (ln.tien_thue if ln.tien_thue is not None else round((ln.thanh_tien or 0) * 0.08))
                entries.append((ln.ten_hang or "Hàng hóa, dịch vụ", gt, thue, "Chiết khấu thương mại, ghi giảm" if sign < 0 else ""))
            if not entries:
                gt, thue = sums[RATE]
                entries.append((f"Hàng hóa, dịch vụ theo HĐ {kh} số {so}", gt, thue, "Hóa đơn không có dòng hàng chi tiết; lấy tổng tiền thuế suất 8%"))
            elif RATE in sums:
                lech = round(sums[RATE][0] - sum(e[1] for e in entries))
                if abs(lech) >= 1:
                    entries[-1] = (*entries[-1][:3], (entries[-1][3] + "; " if entries[-1][3] else "")
                                   + f"Tổng dòng hàng 8% lệch {lech:,.0f} so với bảng thuế suất trên hóa đơn")

            for ten, gt, thue, note in entries:
                k = _norm_name(ten)
                it = bucket.setdefault(k, Item(ten=_cut(ten)))
                it.gia_tri += gt
                it.thue += thue
                giam = round(gt * (TS_QUY_DINH - TS_SAU_GIAM) / 100) if direction == "sold" else None
                detail.append({**base, "ten": ten, "gia_tri": round(gt), "thue": round(thue), "giam": giam, "ghi_chu": note})
    return periods, detail, skipped


# ----------------------------------------------------------------- ghi file .xls theo mẫu


def _layout(rs) -> dict[str, int]:
    """Tìm vị trí 2 phần trong sheet mẫu (theo chữ, không cố định số dòng)."""
    pos: dict[str, int] = {}
    for r in range(rs.nrows):
        v = str(rs.cell_value(r, 1)).strip()
        if v.startswith("I. ") and "i_title" not in pos:
            pos["i_title"] = r
        elif v.startswith("II. "):
            pos["ii_title"] = r
        elif v == "(1)":
            pos["ii_data" if "ii_title" in pos else "i_data"] = r + 1
    if len(pos) != 4:
        raise ValueError(f"Mẫu bảng kê không đúng cấu trúc: {TEMPLATE.name}")
    pos["i_rows"] = pos["ii_title"] - pos["i_data"]
    pos["ii_rows"] = rs.nrows - pos["ii_data"]
    return pos


def write_bang_ke(path: Path, mua: list[Item], ban: list[Item], template: Path = TEMPLATE) -> None:
    import xlrd
    from xlutils.filter import BaseFilter, XLRDReader, XLWTWriter, process

    rb = xlrd.open_workbook(str(template), formatting_info=True)
    rs = rb.sheet_by_name(SHEET)
    pos = _layout(rs)
    n1 = max(len(mua), 1)
    delta = n1 - pos["i_rows"]
    skip = set(range(pos["i_data"], pos["ii_title"])) | set(range(pos["ii_data"], rs.nrows))

    class Remap(BaseFilter):
        active = False

        def sheet(self, rdsheet, wtsheet_name):
            self.active = rdsheet.name == SHEET
            self.next.sheet(rdsheet, wtsheet_name)

        def _map(self, r):
            if not self.active:
                return r
            if r in skip:
                return None
            return r + delta if r >= pos["ii_title"] else r

        def row(self, rdrowx, wtrowx):
            w = self._map(rdrowx)
            if w is not None:
                self.next.row(rdrowx, w)

        def cell(self, rdrowx, rdcolx, wtrowx, wtcolx):
            w = self._map(rdrowx)
            if w is not None:
                self.next.cell(rdrowx, rdcolx, w, wtcolx)

    writer = XLWTWriter()
    process(XLRDReader(rb, path.name), Remap(), writer)
    wb, styles = writer.output[0][1], writer.style_list
    ws = wb.get_sheet(rb.sheet_names().index(SHEET))

    def put_rows(start: int, tpl_row: int, rows: list[list[Any]], count: int) -> None:
        height = rs.rowinfo_map.get(tpl_row)
        for i in range(count):
            r = start + i
            vals = rows[i] if i < len(rows) else [i + 1]
            for c in range(rs.ncols):
                v = vals[c - 1] if 1 <= c <= len(vals) else ""
                ws.write(r, c, v, styles[rs.cell_xf_index(tpl_row, c)])
            if height is not None:
                ws.row(r).height = height.height
                ws.row(r).height_mismatch = True

    put_rows(pos["i_data"], pos["i_data"], [[i + 1, it.ten, round(it.gia_tri), round(it.thue)] for i, it in enumerate(mua)], n1)
    n2 = max(len(ban), 1)
    put_rows(pos["ii_data"] + delta, pos["ii_data"],
             [[i + 1, it.ten, round(it.gia_tri), TS_QUY_DINH, TS_SAU_GIAM, round(it.gia_tri * (TS_QUY_DINH - TS_SAU_GIAM) / 100)]
              for i, it in enumerate(ban)], n2)
    wb.save(str(path))


# ----------------------------------------------------------------- xuất


DETAIL_COLUMNS = [
    ("Kỳ", "ky"), ("Chiều", "chieu"), ("Hóa đơn", "hd"), ("Ngày lập", "ngay"), ("MST đối tác", "mst"), ("Tên đối tác", "doi_tac"),
    ("Tên hàng hóa, dịch vụ", "ten"), ("Giá trị chưa thuế", "gia_tri"), ("Thuế GTGT 8%", "thue"),
    ("Thuế GTGT được giảm (bán ra)", "giam"), ("Ghi chú", "ghi_chu"),
]
SUMMARY_COLUMNS = [
    ("Kỳ", "ky"), ("File bảng kê", "file"), ("Số dòng mua vào", "n_mua"), ("Giá trị mua vào 8%", "gt_mua"), ("Thuế mua vào 8%", "thue_mua"),
    ("Số dòng bán ra", "n_ban"), ("Giá trị bán ra 8%", "gt_ban"), ("Thuế GTGT được giảm", "giam"), ("Ghi chú", "ghi_chu"),
]
SKIP_COLUMNS = [("Chiều", "chieu"), ("Hóa đơn", "hd"), ("Ngày lập", "ngay"), ("Đối tác", "doi_tac"), ("Lý do không đưa vào", "ly_do")]


def _sheet(wb, title: str, columns, rows) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in rows:
        ws.append([r.get(k) for _, k in columns])
    for i, (name, key) in enumerate(columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = {"ten": 50, "doi_tac": 40, "ghi_chu": 50, "ly_do": 40, "file": 40}.get(key, max(12, len(name) + 2))
        for cell in ws[letter][1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy"
            elif isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
    ws.freeze_panes = "A2"


def export_bang_ke(output_dir: str | Path, ky: str = "quy") -> BangKeResult:
    from openpyxl import Workbook

    out_dir = Path(output_dir)
    if not (out_dir / "xml").is_dir():
        raise FileNotFoundError(f"Không thấy hóa đơn đã kéo trong {out_dir / 'xml'}. Hãy kéo hóa đơn (có tải XML) trước.")
    ky = "thang" if str(ky).strip().lower() in ("thang", "tháng", "t", "month") else "quy"
    folder = out_dir / FOLDER
    res = BangKeResult(folder=str(folder))
    periods, detail, skipped = collect(out_dir, ky)
    if not periods:
        log.info("[Bảng kê giảm thuế] Không có hóa đơn thuế suất 8%% trong %s", out_dir)
        return res
    folder.mkdir(parents=True, exist_ok=True)

    summary = []
    for key in sorted(periods):
        per = periods[key]
        notes = []
        mua, ban = [], []
        for direction, target in (("purchase", mua), ("sold", ban)):
            for it in per[direction].values():
                if round(it.gia_tri) <= 0:
                    notes.append(f"{'Mua vào' if direction == 'purchase' else 'Bán ra'} '{it.ten}': tổng giá trị {round(it.gia_tri):,} không dương, không đưa vào bảng kê")
                    continue
                target.append(it)
        name = f"BangKe01_GiamThue_{per['file']}.xls"
        write_bang_ke(folder / name, mua, ban)
        res.files.append(str(folder / name))
        info = {
            "ky": per["label"], "file": name, "n_mua": len(mua), "gt_mua": round(sum(i.gia_tri for i in mua)),
            "thue_mua": round(sum(i.thue for i in mua)), "n_ban": len(ban), "gt_ban": round(sum(i.gia_tri for i in ban)),
            "giam": sum(round(i.gia_tri * (TS_QUY_DINH - TS_SAU_GIAM) / 100) for i in ban), "ghi_chu": "; ".join(notes),
        }
        summary.append(info)
        res.warnings += [f"{per['label']}: {n}" for n in notes]
    res.periods = summary

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "TongHop", SUMMARY_COLUMNS, summary)
    _sheet(wb, "ChiTiet", DETAIL_COLUMNS, detail)
    _sheet(wb, "KhongDuaVao", SKIP_COLUMNS, skipped)
    res.detail_path = str(folder / DETAIL_FILE)
    wb.save(res.detail_path)
    log.info("[Bảng kê giảm thuế] %d file (%s) -> %s", len(res.files), res.summary(), folder)
    return res
