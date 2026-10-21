"""Xuất kết quả ra Excel (openpyxl): sheet HoaDon, ChiTiet, Loi."""

from __future__ import annotations

import re
from datetime import datetime
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .client import TTHAI_LABELS, TTXLY_LABELS, InvoiceRef
from .thuesuat import BUCKETS, TAXED_BUCKETS, UNKNOWN, breakdown
from .xmlparse import ParsedInvoice, parse_date

SLUG = {"0%": "0", "5%": "5", "8%": "8", "10%": "10", "KCT": "kct", "KKKNT": "kkknt", "Khác": "khac", UNKNOWN: "chuaro"}
# Cột tiền theo từng mức thuế suất (đặt ngay sau các cột tổng tiền của hóa đơn).
RATE_COLUMNS: list[tuple[str, str]] = (
    [(f"Chưa thuế {b}" if b != "Khác" else "Chưa thuế suất khác", f"cth_{SLUG[b]}") for b in BUCKETS]
    + [("Chưa thuế chưa rõ thuế suất", "cth_chuaro")]
    + [(f"Thuế GTGT {b}" if b != "Khác" else "Thuế GTGT suất khác", f"thue_{SLUG[b]}") for b in TAXED_BUCKETS]
    + [("Thuế GTGT chưa rõ thuế suất", "thue_chuaro"), ("Lệch tiền hàng so với tổng HĐ", "cth_lech")]
)
# Hóa đơn bị thay thế (4), bị hủy (6) không tính vào tổng hợp theo thuế suất.
EXCLUDED_TTHAI = {4, 6}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")

INVOICE_COLUMNS: list[tuple[str, str]] = [
    ("Chiều", "chieu"),
    ("Nguồn", "nguon"),
    ("Ký hiệu mẫu số", "khmshdon"),
    ("Ký hiệu hóa đơn", "khhdon"),
    ("Số hóa đơn", "shdon"),
    ("Ngày lập", "ngay_lap"),
    ("MST người bán", "nbmst"),
    ("Tên người bán", "nbten"),
    ("Địa chỉ người bán", "nbdchi"),
    ("MST người mua", "nmmst"),
    ("Tên người mua", "nmten"),
    ("Địa chỉ người mua", "nmdchi"),
    ("Tổng tiền chưa thuế", "tgtcthue"),
    ("Tổng tiền thuế", "tgtthue"),
    ("Tổng tiền chiết khấu", "ttcktmai"),
    ("Tổng tiền thanh toán", "tgtttbso"),
    *RATE_COLUMNS,
    ("Đơn vị tiền tệ", "dvtte"),
    ("Tỷ giá", "tgia"),
    ("Trạng thái hóa đơn", "tthai_label"),
    ("Kết quả kiểm tra", "ttxly_label"),
    ("Mã CQT", "mhdon"),
    ("Ngày ký", "nky"),
    ("Ngày cấp mã", "ncma"),
    ("Hóa đơn gốc (mẫu số/ký hiệu/số)", "hd_goc"),
    ("Ngày HĐ gốc", "tdlhdgoc"),
    ("Ghi chú HĐ gốc", "gchdgoc"),
    ("Hình thức thanh toán", "htttoan"),
    ("MST TCGP", "msttcgp"),
    ("File PDF", "pdf_file"),
    ("File XML/JSON", "xml_file"),
    ("Ghi chú tải XML", "xml_error"),
    ("Nguồn số liệu thuế suất", "ts_nguon"),
    ("Quý", "ky_quy"),
    ("Tháng", "ky_thang"),
    ("Tính vào tổng hợp", "tinh_th"),
]

LINE_COLUMNS: list[tuple[str, str]] = [
    ("Chiều", "chieu"),
    ("Ký hiệu mẫu số", "khmshdon"),
    ("Ký hiệu hóa đơn", "khhdon"),
    ("Số hóa đơn", "shdon"),
    ("Ngày lập", "ngay_lap"),
    ("MST người bán", "nbmst"),
    ("Tên người bán", "nbten"),
    ("MST người mua", "nmmst"),
    ("Tên người mua", "nmten"),
    ("STT", "stt"),
    ("Tính chất", "tchat"),
    ("Mã hàng", "ma_hang"),
    ("Tên hàng hóa, dịch vụ", "ten_hang"),
    ("ĐVT", "dvt"),
    ("Số lượng", "so_luong"),
    ("Đơn giá", "don_gia"),
    ("Tỷ lệ CK", "tl_ck"),
    ("Số tiền CK", "st_ck"),
    ("Thành tiền", "thanh_tien"),
    ("Thuế suất", "thue_suat"),
    ("Tiền thuế", "tien_thue"),
    ("Thành tiền có thuế", "thanh_tien_co_thue"),
]

ERROR_COLUMNS = ["Chiều", "Nguồn", "MST người bán", "Mẫu số", "Ký hiệu", "Số HĐ", "Bước", "Lỗi"]

TCHAT_LABELS = {"1": "Hàng hóa, dịch vụ", "2": "Khuyến mại", "3": "Chiết khấu thương mại", "4": "Ghi chú/diễn giải"}


def _label(mapping: dict[int, str], value: Any) -> str:
    try:
        n = int(str(value))
    except (TypeError, ValueError):
        return str(value or "")
    return mapping.get(n, str(value))


def _num(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return value


def _date(value: Any) -> Any:
    if isinstance(value, datetime):
        return value
    dt = parse_date(str(value)) if value else None
    return dt or (value or None)


def invoice_row(ref: InvoiceRef, parsed: ParsedInvoice | None = None, xml_file: str = "", xml_error: str = "") -> dict[str, Any]:
    r = ref.raw
    hd_goc = ""
    if r.get("khmshdgoc") or r.get("shdgoc"):
        hd_goc = f"{r.get('khmshdgoc') or ''}{r.get('khhdgoc') or ''} / {r.get('shdgoc') or ''}"
    row: dict[str, Any] = {
        "chieu": "Mua vào" if ref.direction == "purchase" else "Bán ra",
        "nguon": "Máy tính tiền" if ref.family == "sco-query" else "Hóa đơn điện tử",
        "khmshdon": ref.khmshdon,
        "khhdon": ref.khhdon,
        "shdon": _num(ref.shdon),
        "ngay_lap": _date(r.get("tdlap")),
        "nbmst": ref.nbmst,
        "nbten": r.get("nbten", ""),
        "nbdchi": r.get("nbdchi", ""),
        "nmmst": r.get("nmmst", ""),
        "nmten": r.get("nmten") or r.get("nmtnmua") or "",
        "nmdchi": r.get("nmdchi", ""),
        "tgtcthue": _num(r.get("tgtcthue")),
        "tgtthue": _num(r.get("tgtthue")),
        "ttcktmai": _num(r.get("ttcktmai")),
        "tgtttbso": _num(r.get("tgtttbso")),
        "dvtte": r.get("dvtte", ""),
        "tgia": _num(r.get("tgia")),
        "tthai_label": _label(TTHAI_LABELS, r.get("tthai")),
        "ttxly_label": _label(TTXLY_LABELS, r.get("ttxly")),
        "mhdon": r.get("mhdon", ""),
        "nky": _date(r.get("nky")),
        "ncma": _date(r.get("ncma")),
        "hd_goc": hd_goc,
        "tdlhdgoc": _date(r.get("tdlhdgoc")),
        "gchdgoc": r.get("gchdgoc", ""),
        "htttoan": r.get("thtttoan") or r.get("htttoan", ""),
        "msttcgp": r.get("msttcgp", ""),
        "xml_file": xml_file,
        "xml_error": xml_error,
    }
    if parsed is not None:
        # Ưu tiên số liệu trong XML gốc nếu danh sách thiếu.
        fill = {
            "nbten": parsed.nb_ten, "nbdchi": parsed.nb_dchi, "nmmst": parsed.nm_mst, "nmten": parsed.nm_ten or parsed.nm_hvtn,
            "nmdchi": parsed.nm_dchi, "tgtcthue": parsed.tong_tien_chua_thue, "tgtthue": parsed.tong_tien_thue,
            "ttcktmai": parsed.tong_ck, "tgtttbso": parsed.tong_tien_tt, "dvtte": parsed.dvtte, "tgia": parsed.ty_gia,
            "mhdon": parsed.mccqt, "nky": parsed.ngay_ky_nb, "ncma": parsed.ngay_ky_cqt, "htttoan": parsed.httt,
            "msttcgp": parsed.mst_tcgp, "ngay_lap": parsed.ngay_lap,
        }
        for k, v in fill.items():
            if (row.get(k) in (None, "")) and v not in (None, ""):
                row[k] = v
    _add_rate_breakdown(row, r, parsed)
    return row


def _round(v: float) -> float | int | None:
    v = round(v, 2)
    if abs(v) < 0.005:
        return None
    return int(v) if float(v).is_integer() else v


def _add_rate_breakdown(row: dict[str, Any], raw: dict[str, Any], parsed: ParsedInvoice | None) -> None:
    sums, src = breakdown(raw, parsed)
    total = row.get("tgtcthue")
    if not sums:
        # Không có bảng thuế suất lẫn dòng hàng: để cả hóa đơn vào "chưa rõ" cho tổng vẫn khớp.
        sums = {UNKNOWN: [float(total or 0), float(row.get("tgtthue") or 0)]}
        src = "Không có số liệu theo thuế suất"
    for b in (*BUCKETS, UNKNOWN):
        amt, tax = sums.get(b, [0.0, 0.0])
        row[f"cth_{SLUG[b]}"] = _round(amt)
        if b in TAXED_BUCKETS or b == UNKNOWN:
            row[f"thue_{SLUG[b]}"] = _round(tax)
    if isinstance(total, (int, float)):
        row["cth_lech"] = _round(total - sum(v[0] for v in sums.values()))
    row["ts_nguon"] = src
    d = row.get("ngay_lap")
    if isinstance(d, datetime):
        row["ky_quy"] = f"Quý {(d.month - 1) // 3 + 1}/{d.year}"
        row["ky_thang"] = f"Tháng {d.month:02d}/{d.year}"
    try:
        tthai = int(raw.get("tthai")) if raw.get("tthai") not in (None, "") else 0
    except (TypeError, ValueError):
        tthai = 0
    row["tinh_th"] = "Không" if tthai in EXCLUDED_TTHAI else "Có"


def line_rows(ref: InvoiceRef, parsed: ParsedInvoice) -> list[dict[str, Any]]:
    base = {
        "chieu": "Mua vào" if ref.direction == "purchase" else "Bán ra",
        "khmshdon": ref.khmshdon or parsed.khmshdon,
        "khhdon": ref.khhdon or parsed.khhdon,
        "shdon": _num(ref.shdon or parsed.shdon),
        "ngay_lap": parsed.ngay_lap or _date(ref.raw.get("tdlap")),
        "nbmst": ref.nbmst or parsed.nb_mst,
        "nbten": parsed.nb_ten or ref.raw.get("nbten", ""),
        "nmmst": parsed.nm_mst or ref.raw.get("nmmst", ""),
        "nmten": parsed.nm_ten or parsed.nm_hvtn or ref.raw.get("nmten") or ref.raw.get("nmtnmua") or "",
    }
    rows = []
    for ln in parsed.lines:
        rows.append(
            {
                **base,
                "stt": _num(ln.stt),
                "tchat": TCHAT_LABELS.get(ln.tchat, ln.tchat),
                "ma_hang": ln.ma_hang,
                "ten_hang": ln.ten_hang,
                "dvt": ln.dvt,
                "so_luong": ln.so_luong,
                "don_gia": ln.don_gia,
                "tl_ck": ln.tl_ck,
                "st_ck": ln.st_ck,
                "thanh_tien": ln.thanh_tien,
                "thue_suat": ln.thue_suat,
                "tien_thue": ln.tien_thue,
                "thanh_tien_co_thue": ln.thanh_tien_co_thue,
            }
        )
    return rows


AMOUNT_KEYS = {"tgtcthue", "tgtthue", "ttcktmai", "tgtttbso", "thanh_tien", "tien_thue", "thanh_tien_co_thue", "don_gia", "st_ck"}


def _amount_format(v: Any) -> str:
    """Số nguyên hiện 1.234.567 (không có dấu thập phân thừa); số lẻ hiện tối đa 4 chữ số thập phân."""
    if isinstance(v, float) and not v.is_integer():
        return "#,##0.####"
    return "#,##0"


def _write_sheet(ws, columns: list[tuple[str, str]], rows: Iterable[dict[str, Any]]) -> None:
    ws.append([h for h, _ in columns])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    count = 0
    for row in rows:
        ws.append([row.get(key) for _, key in columns])
        count += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(count + 1, 1)}"
    for idx, (header, key) in enumerate(columns, start=1):
        width = max(10, min(45, len(header) + 2))
        if key in ("nbten", "nmten", "ten_hang", "nbdchi", "nmdchi"):
            width = 40
        ws.column_dimensions[get_column_letter(idx)].width = width
        if key in ("ngay_lap", "nky", "ncma", "tdlhdgoc"):
            for cell in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell[0].number_format = "dd/mm/yyyy"
        if key in AMOUNT_KEYS or key.startswith(("cth_", "thue_")):
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = _amount_format(cell.value)


def write_workbook(
    path: str,
    invoices: list[dict[str, Any]],
    lines: list[dict[str, Any]],
    errors: list[dict[str, Any]],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "HoaDon"
    invoices = list(invoices)
    _write_sheet(ws, INVOICE_COLUMNS, invoices)
    write_rate_summary(wb, invoices)
    _write_sheet(wb.create_sheet("ChiTiet"), LINE_COLUMNS, lines)
    err_cols = [(h, h) for h in ERROR_COLUMNS]
    _write_sheet(wb.create_sheet("Loi"), err_cols, errors)
    wb.save(path)


def _period_key(label: str) -> tuple[int, int]:
    nums = [int(x) for x in re.findall(r"\d+", label or "")]
    return (nums[-1], nums[0]) if len(nums) >= 2 else (0, 0)


def write_rate_summary(wb, invoices: list[dict[str, Any]], sheet_title: str = "TongHopThueSuat") -> None:
    """Sheet tổng hợp tiền chưa thuế / tiền thuế theo thuế suất, theo quý và theo tháng.

    Mọi con số là công thức SUMIFS/COUNTIFS lấy từ sheet HoaDon, chỉ cộng hóa đơn có
    "Tính vào tổng hợp" = Có, nên sửa số liệu ở HoaDon thì bảng này tự cập nhật.
    """
    ws = wb.create_sheet(sheet_title, 1)
    col = {key: get_column_letter(i) for i, (_, key) in enumerate(INVOICE_COLUMNS, start=1)}
    last = max(len(invoices) + 1, 2)

    def rng(key: str) -> str:
        return f"HoaDon!${col[key]}$2:${col[key]}${last}"

    amt_cols = [(f"Chưa thuế {b}" if b != "Khác" else "Chưa thuế suất khác", f"cth_{SLUG[b]}") for b in BUCKETS]
    amt_cols.append(("Chưa rõ thuế suất", "cth_chuaro"))
    tax_cols = [(f"Thuế GTGT {b}" if b != "Khác" else "Thuế GTGT suất khác", f"thue_{SLUG[b]}") for b in TAXED_BUCKETS]
    tax_cols.append(("Thuế chưa rõ thuế suất", "thue_chuaro"))
    headers = ["Chiều", "Kỳ", "Số hóa đơn"] + [h for h, _ in amt_cols] + ["Cộng tiền chưa thuế"] + [h for h, _ in tax_cols] + ["Cộng tiền thuế"]
    n_amt, n_tax = len(amt_cols), len(tax_cols)
    c_amt0 = 4
    c_amt_sum = c_amt0 + n_amt
    c_tax0 = c_amt_sum + 1
    c_tax_sum = c_tax0 + n_tax
    L = get_column_letter

    title_font = Font(bold=True, size=13, color="1F4E78")
    bold = Font(bold=True)
    total_fill = PatternFill("solid", fgColor="DDEBF7")
    ws["A1"] = "TỔNG HỢP TIỀN HÀNG CHƯA THUẾ VÀ TIỀN THUẾ GTGT THEO TỪNG MỨC THUẾ SUẤT"
    ws["A1"].font = title_font
    ws["A2"] = ("Số liệu là công thức lấy từ sheet HoaDon. Không tính hóa đơn bị hủy, bị thay thế "
                "(cột \"Tính vào tổng hợp\" = Không).")
    ws["A2"].font = Font(italic=True, color="595959")
    ws["A3"] = "Số hóa đơn không tính vào tổng hợp:"
    ws["E3"] = f'=COUNTIFS({rng("tinh_th")},"Không")'
    ws["E3"].font = bold

    directions = [d for d in ("Mua vào", "Bán ra") if any(r.get("chieu") == d for r in invoices)]
    row = 5
    for block_title, pkey in (("THEO QUÝ", "ky_quy"), ("THEO THÁNG", "ky_thang")):
        ws.cell(row, 1, block_title).font = Font(bold=True, color="1F4E78")
        row += 1
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row, i, h)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 32
        row += 1
        for d in directions:
            periods = sorted({r.get(pkey) for r in invoices if r.get("chieu") == d and r.get(pkey)}, key=_period_key)
            first = row
            for p in periods:
                ws.cell(row, 1, d)
                ws.cell(row, 2, p)
                crit = f'{rng("chieu")},$A{row},{rng(pkey)},$B{row},{rng("tinh_th")},"Có"'
                ws.cell(row, 3, f"=COUNTIFS({crit})")
                for j, (_, key) in enumerate(amt_cols):
                    ws.cell(row, c_amt0 + j, f"=SUMIFS({rng(key)},{crit})")
                ws.cell(row, c_amt_sum, f"=SUM({L(c_amt0)}{row}:{L(c_amt_sum - 1)}{row})")
                for j, (_, key) in enumerate(tax_cols):
                    ws.cell(row, c_tax0 + j, f"=SUMIFS({rng(key)},{crit})")
                ws.cell(row, c_tax_sum, f"=SUM({L(c_tax0)}{row}:{L(c_tax_sum - 1)}{row})")
                row += 1
            # dòng cộng của chiều
            ws.cell(row, 1, d)
            ws.cell(row, 2, "Cộng")
            for c in range(3, c_tax_sum + 1):
                ws.cell(row, c, f"=SUM({L(c)}{first}:{L(c)}{row - 1})" if row > first else 0)
            for c in range(1, c_tax_sum + 1):
                ws.cell(row, c).font = bold
                ws.cell(row, c).fill = total_fill
            row += 1
        row += 1

    for r in ws.iter_rows(min_row=7, max_row=row, min_col=3, max_col=c_tax_sum):
        for cell in r:
            cell.number_format = '#,##0;-#,##0;"-"'  # số 0 hiện dấu "-"
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 10
    for c in range(c_amt0, c_tax_sum + 1):
        ws.column_dimensions[L(c)].width = 16
    ws.freeze_panes = "C7"
