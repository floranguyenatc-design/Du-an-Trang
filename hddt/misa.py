"""Chuẩn bị dữ liệu hóa đơn mua vào đã kéo về để nhập vào phần mềm MISA SME.

Kết quả trong thư mục ``misa``:

- ``Mua_hang_khong_qua_kho_VND.xlsx``: điền đúng mẫu nhập khẩu "Chứng từ mua hàng không
  qua kho" (VND) của MISA SME.NET, giữ nguyên 39 cột, thứ tự cột và ghi chú gợi ý của
  MISA. Mỗi dòng hàng hóa một dòng; các dòng cùng "Số chứng từ" là một chứng từ.
  Các cột mã hóa theo quy định của mẫu: Hình thức mua hàng 0 = trong nước, Phương thức
  thanh toán 0 = chưa thanh toán, Nhận kèm hóa đơn 1 = có, % thuế GTGT 0/5/8/10/KCT/
  KKKNT/KHAC (KHAC kèm Tỷ lệ tính thuế).
- ``MISA_DanhMuc_va_KiemTra.xlsx``: danh mục Nhà cung cấp, Vật tư hàng hóa (mã NCC và mã
  hàng phải có trong MISA trước khi nhập chứng từ), sheet KiemTra đối chiếu tổng dòng
  hàng với tổng hóa đơn và ghi chú cần xem lại, sheet BoQua.
- ``xml_mua_vao``: XML gốc để dùng chức năng đọc hóa đơn điện tử của MISA (tùy chọn).

Hóa đơn bị hủy (tthai 6) và bị thay thế (tthai 4) không đưa vào.
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
import shutil
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .client import TTHAI_LABELS
from .detail import parsed_from_detail
from .xmlparse import ParsedInvoice, parse_invoice_file

log = logging.getLogger("hddt")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
REQUIRED_FILL = PatternFill("solid", fgColor="B71C1C")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# Trạng thái không hạch toán: 4 = bị thay thế, 6 = bị hủy.
SKIP_TTHAI = {4: "Hóa đơn đã bị thay thế (hạch toán theo hóa đơn thay thế)", 6: "Hóa đơn đã bị hủy"}

TEMPLATE_PATH = Path(__file__).parent / "templates" / "misa_mua_hang_khong_qua_kho_vnd.xlsx"
TEMPLATE_OUT_NAME = "Mua_hang_khong_qua_kho_VND.xlsx"

# Cột của mẫu MISA "Chứng từ mua hàng không qua kho" (đúng thứ tự A..AM).
MISA_COLUMNS: list[tuple[str, str]] = [
    ("A", "Hiển thị trên sổ"), ("B", "Hình thức mua hàng"), ("C", "Phương thức thanh toán"),
    ("D", "Nhận kèm hóa đơn"), ("E", "Ngày hạch toán (*)"), ("F", "Ngày chứng từ (*)"),
    ("G", "Số chứng từ (*)"), ("H", "Mẫu số HĐ"), ("I", "Ký hiệu HĐ"), ("J", "Số hóa đơn"),
    ("K", "Ngày hóa đơn"), ("L", "Mã nhà cung cấp"), ("M", "Tên nhà cung cấp"), ("N", "Diễn giải"),
    ("O", "NV mua hàng"), ("P", "Mã hàng (*)"), ("Q", "Tên hàng"), ("R", "TK chi phí (*)"),
    ("S", "TK công nợ/TK tiền (*)"), ("T", "ĐVT"), ("U", "Số lượng"), ("V", "Đơn giá"),
    ("W", "Thành tiền"), ("X", "Tỷ lệ CK"), ("Y", "Tiền chiết khấu"), ("Z", "Chi phí mua hàng"),
    ("AA", "% thuế GTGT"), ("AB", "Tỷ lệ tính thuế (Thuế suất KHAC)"), ("AC", "Tiền thuế GTGT"),
    ("AD", "TKĐƯ thuế GTGT"), ("AE", "TK thuế GTGT"), ("AF", "Nhóm HHDV mua vào"),
    ("AG", "Giá tính thuế NK"), ("AH", "% thuế NK"), ("AI", "Tiền thuế NK"), ("AJ", "TK thuế NK"),
    ("AK", "% thuế TTĐB"), ("AL", "Tiền thuế TTĐB"), ("AM", "TK thuế TTĐB"),
]
DATE_COLS = ("E", "F", "K")
TEXT_COLS = ("G", "H", "I", "J", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "AA", "AD", "AE", "AF")
NUM_COLS = ("U", "V", "W", "X", "Y", "AB", "AC")
# Giới hạn độ dài theo ghi chú của mẫu MISA.
MAX_LEN = {"G": 20, "H": 25, "I": 20, "J": 25, "M": 128, "N": 255, "Q": 255, "R": 20, "S": 25}

# Phương thức thanh toán (cột C) -> tài khoản đối ứng mặc định (cột S).
PAYMENT_ACCOUNT = {"0": None, "1": "1111", "2": "1121", "3": "1121", "4": "1121"}

SUPPLIER_COLUMNS: list[tuple[str, str, bool]] = [
    ("Mã nhà cung cấp", "ma", True),
    ("Tên nhà cung cấp", "ten", True),
    ("Địa chỉ", "dia_chi", False),
    ("Mã số thuế", "mst", False),
    ("Điện thoại", "dien_thoai", False),
    ("Tài khoản ngân hàng", "stk", False),
    ("Tổ chức/Cá nhân", "loai", False),
]

ITEM_COLUMNS: list[tuple[str, str, bool]] = [
    ("Mã", "ma", True),
    ("Tên", "ten", True),
    ("Tính chất", "tinh_chat", True),
    ("Đơn vị tính chính", "dvt", False),
    ("Thuế suất GTGT (%)", "thue_suat", False),
    ("TK kho", "tk_kho", False),
    ("TK chi phí", "tk_chi_phi", False),
    ("Mã hàng trên hóa đơn", "ma_goc", False),
]


@dataclass
class MisaSettings:
    tk_chi_phi: str = "642"
    tk_cong_no: str = "331"
    tk_thue: str = "1331"
    nhom_hhdv: str = "1"  # mã nhóm HHDV mua vào trong danh mục MISA
    phuong_thuc_tt: str = "0"  # 0 chưa thanh toán, 1 tiền mặt, 2 ủy nhiệm chi, 3 séc CK, 4 séc TM
    so_ct_prefix: str = "MH"


@dataclass
class MisaResult:
    folder: str = ""
    excel_path: str = ""  # file theo mẫu MISA
    catalog_path: str = ""  # danh mục + kiểm tra
    xml_dir: str = ""
    vouchers: int = 0
    lines: int = 0
    xml_files: int = 0
    no_xml: int = 0
    skipped: int = 0
    suppliers: int = 0
    items: int = 0
    warnings: list[str] = field(default_factory=list)


def misa_vat_rate(text: str) -> tuple[str, float | None]:
    """Chuyển thuế suất trên hóa đơn sang mã cột "% thuế GTGT" của MISA.

    Trả về (mã, tỷ lệ cho thuế suất KHAC). Mã hợp lệ: 0, 5, 8, 10, KCT, KKKNT, KHAC.
    """
    t = (text or "").strip().upper().replace(" ", "")
    if not t:
        return "", None
    if t in ("KCT", "KKKNT"):
        return t, None
    if t.startswith("KHAC"):
        m = re.search(r"(\d+(?:[.,]\d+)?)", t)
        return "KHAC", (float(m.group(1).replace(",", ".")) if m else None)
    m = re.fullmatch(r"(\d+(?:[.,]\d+)?)%?", t)
    if not m:
        return "", None
    v = float(m.group(1).replace(",", "."))
    if 0 < v < 1 and "%" not in t:
        v *= 100
    if v in (0, 5, 8, 10):
        return str(int(v)), None
    return "KHAC", v


def _vnd(v: float | None, decimals: int = 0) -> float | int | None:
    if v is None:
        return None
    r = round(v, decimals)
    return int(r) if float(r).is_integer() else r


def _cut(text: str, n: int) -> str:
    text = text or ""
    return text if len(text) <= n else text[: n - 1] + "…"


def _ascii_slug(text: str) -> str:
    t = unicodedata.normalize("NFD", text or "")
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = t.replace("đ", "d").replace("Đ", "D")
    t = re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_").upper()
    return t


def _norm_name(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip()).lower()


class ItemCatalog:
    """Gán mã hàng ổn định: dùng mã trên hóa đơn nếu có, không thì tạo từ tên hàng."""

    def __init__(self) -> None:
        self.by_key: dict[tuple[str, str, str], dict[str, Any]] = {}
        self.used: dict[str, tuple[str, str, str]] = {}

    def code_for(self, ma_goc: str, ten: str, dvt: str, thue_suat: str, s: MisaSettings) -> str:
        # Khóa theo tên + ĐVT: có người bán dùng một mã (vd SP.0000) cho nhiều mặt hàng.
        key = (_norm_name(ten), _norm_name(dvt))
        if key in self.by_key:
            return self.by_key[key]["ma"]
        goc = (ma_goc or "").strip()
        base = _ascii_slug(goc)[:20] if goc else (_ascii_slug(ten)[:18] or "HH")
        code = base
        if code in self.used and self.used[code] != key:
            h = hashlib.md5("|".join(key).encode("utf-8")).hexdigest()[:4].upper()
            code = f"{base[:20]}_{h}"
        self.used[code] = key
        name = _norm_name(ten)
        dich_vu = any(w in name for w in ("cước", "dịch vụ", "phí ", "phí", "thuê", "vận chuyển", "sửa chữa", "bảo trì"))
        self.by_key[key] = {
            "ma": code,
            "ten": ten,
            "tinh_chat": "Dịch vụ" if dich_vu else "Vật tư hàng hóa",
            "dvt": dvt,
            "thue_suat": thue_suat,
            "tk_kho": "" if dich_vu else "156",
            "tk_chi_phi": s.tk_chi_phi,
            "ma_goc": goc,
        }
        return code


def _load_raw_index(out_dir: Path) -> dict[str, dict[str, Any]]:
    """Ghép các file danh_sach_*.json: khóa = file_stem của hóa đơn."""
    from .client import InvoiceRef

    index: dict[str, dict[str, Any]] = {}
    for f in sorted(out_dir.glob("danh_sach_*.json")):
        try:
            items = json.loads(f.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        for item in items if isinstance(items, list) else []:
            if not isinstance(item, dict):
                continue
            for fam in ("query", "sco-query"):
                ref = InvoiceRef.from_json("purchase", fam, item)
                index[ref.file_stem] = item
    return index


def _collect(out_dir: Path, raw_index: dict[str, dict[str, Any]]) -> list[tuple[str, Path, ParsedInvoice, bool]]:
    """Trả về [(stem, file nguồn, ParsedInvoice, có XML gốc)] cho hóa đơn mua vào."""
    src = out_dir / "xml" / "purchase"
    found: dict[str, tuple[str, Path, ParsedInvoice, bool]] = {}
    if not src.is_dir():
        return []
    for xml in sorted(src.glob("*.xml")):
        stem = xml.stem
        # Gói có nhiều XML được lưu là <stem>_1.xml, <stem>_2.xml: lấy file đầu tiên.
        base = re.sub(r"_\d+$", "", stem)
        if stem not in raw_index and base in raw_index:
            stem = base
        if stem in found:
            continue
        try:
            found[stem] = (stem, xml, parse_invoice_file(str(xml)), True)
        except Exception as exc:  # noqa: BLE001
            log.warning("[MISA] Bỏ qua %s: đọc XML lỗi: %s", xml.name, exc)
    for js in sorted(src.glob("*.json")):
        if js.stem in found:
            continue
        try:
            found[js.stem] = (js.stem, js, parsed_from_detail(json.loads(js.read_text(encoding="utf-8"))), False)
        except Exception as exc:  # noqa: BLE001
            log.warning("[MISA] Bỏ qua %s: đọc chi tiết lỗi: %s", js.name, exc)
    return list(found.values())


def _write(ws, columns: list[tuple[str, str, bool]], rows: list[dict[str, Any]]) -> None:
    ws.append([c[0] for c in columns])
    for i, (_, _, req) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill = REQUIRED_FILL if req else HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r.get(c[1]) for c in columns])
    ws.freeze_panes = "A2"
    for i, (name, key, _) in enumerate(columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = 45 if key in ("hd", "dia_chi", "ten", "ghi_chu", "ly_do") else max(11, len(name) + 3)
        if key.startswith("ngay"):
            for (c,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                c.number_format = "dd/mm/yyyy"
        if key in ("tt_dong", "tt_hd", "lech_tt", "thue_dong", "thue_hd", "lech_thue", "tong_tt"):
            for (c,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                c.number_format = "#,##0.##"
        if key in ("ma_ncc", "mst", "ma", "so_ct", "so_hd", "mau_so", "tk_no", "tk_co", "tk_thue", "ma_hang"):
            for (c,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                c.number_format = "@"


def _fill_template(path: Path, rows: list[dict[str, Any]]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    for i, row in enumerate(rows, start=2):
        for col, _ in MISA_COLUMNS:
            v = row.get(col)
            if v in (None, ""):
                continue
            if col in MAX_LEN and isinstance(v, str):
                v = _cut(v, MAX_LEN[col])
            cell = ws[f"{col}{i}"]
            cell.value = v
            if col in DATE_COLS:
                cell.number_format = "dd/mm/yyyy"
            elif col in TEXT_COLS:
                cell.number_format = "@"
            elif col in NUM_COLS:
                cell.number_format = "#,##0.####"
    wb.save(path)


def write_misa_error(output_dir: str | Path, exc: BaseException) -> str:
    """Ghi lỗi vào misa/LOI_XUAT_MISA.txt để người dùng thấy lý do ngay trong thư mục misa."""
    import traceback

    misa_dir = Path(output_dir) / "misa"
    misa_dir.mkdir(parents=True, exist_ok=True)
    text = (
        f"{datetime.now():%d/%m/%Y %H:%M:%S} - Không tạo được file nhập MISA.\n"
        f"Lý do: {exc}\n\nChi tiết kỹ thuật (gửi cho người hỗ trợ):\n"
        + "".join(traceback.format_exception(type(exc), exc, exc.__traceback__))
    )
    (misa_dir / "LOI_XUAT_MISA.txt").write_text(text, encoding="utf-8")
    return str(misa_dir)


def export_misa(output_dir: str | Path, settings: MisaSettings | None = None) -> MisaResult:
    s = settings or MisaSettings()
    out_dir = Path(output_dir)
    misa_dir = out_dir / "misa"
    xml_out = misa_dir / "xml_mua_vao"
    if xml_out.exists():
        shutil.rmtree(xml_out)
    xml_out.mkdir(parents=True, exist_ok=True)
    res = MisaResult(folder=str(misa_dir), xml_dir=str(xml_out))

    raw_index = _load_raw_index(out_dir)
    invoices = _collect(out_dir, raw_index)
    if not invoices:
        raise FileNotFoundError(
            f"Không thấy hóa đơn mua vào trong {out_dir / 'xml' / 'purchase'}. Hãy kéo hóa đơn (có tải XML) trước."
        )
    invoices.sort(key=lambda t: (t[2].ngay_lap or datetime.min, t[2].nb_mst, t[2].khhdon, t[2].shdon))

    pay = s.phuong_thuc_tt if s.phuong_thuc_tt in PAYMENT_ACCOUNT else "0"
    tk_doi_ung = PAYMENT_ACCOUNT[pay] or s.tk_cong_no

    catalog = ItemCatalog()
    suppliers: dict[str, dict[str, Any]] = {}
    rows: list[dict[str, Any]] = []
    checks: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    used_so_ct: set[str] = set()

    for stem, path, inv, has_xml in invoices:
        raw = raw_index.get(stem, {})
        try:
            tthai = int(raw.get("tthai")) if raw.get("tthai") not in (None, "") else 0
        except (TypeError, ValueError):
            tthai = 0
        label = f"{inv.khmshdon}{inv.khhdon} số {inv.shdon} - {inv.nb_ten}"
        if tthai in SKIP_TTHAI:
            skipped.append({"hd": label, "mst": inv.nb_mst, "ly_do": SKIP_TTHAI[tthai], "file": path.name})
            continue

        ma_ncc = inv.nb_mst or _ascii_slug(inv.nb_ten)[:20] or "NCC"
        suppliers.setdefault(ma_ncc, {
            "ma": ma_ncc, "ten": inv.nb_ten, "dia_chi": inv.nb_dchi, "mst": inv.nb_mst,
            "dien_thoai": inv.nb_sdt, "stk": inv.nb_stk, "loai": "Tổ chức",
        })

        base_ct = f"{s.so_ct_prefix}{_ascii_slug(inv.khhdon)}{inv.shdon}"
        so_ct = base_ct[:20]
        n = 1
        while so_ct in used_so_ct:
            n += 1
            so_ct = base_ct[: 20 - len(str(n)) - 1] + f"_{n}"
        used_so_ct.add(so_ct)

        notes = []
        if not has_xml:
            notes.append("Không có XML gốc; số liệu lấy từ chi tiết trên cổng thuế")
            res.no_xml += 1
        if tthai in (2, 3, 5):
            notes.append(TTHAI_LABELS.get(tthai, "") + " - kiểm tra lại trước khi hạch toán")

        header = {
            "B": 0, "C": int(pay), "D": 1,
            "E": inv.ngay_lap, "F": inv.ngay_lap, "G": so_ct,
            "H": inv.khmshdon, "I": inv.khhdon, "J": inv.shdon, "K": inv.ngay_lap,
            "L": ma_ncc, "M": inv.nb_ten,
            "N": f"Mua hàng của {inv.nb_ten} theo HĐ {inv.khhdon} số {inv.shdon}",
            "R": s.tk_chi_phi, "S": tk_doi_ung,
        }

        def add_line(ma: str, ten: str, dvt: str, sl, dg, tt, tl_ck, tien_ck, ts_text: str, thue) -> None:
            code, khac = misa_vat_rate(ts_text)
            row = {
                **header, "P": ma, "Q": ten, "T": dvt,
                "U": _vnd(sl, 4), "V": _vnd(dg, 4), "W": _vnd(tt, 2),
                "X": _vnd(tl_ck, 4) if tl_ck else None, "Y": _vnd(tien_ck, 2) if tien_ck else None,
                "AA": code, "AB": _vnd(khac, 4) if code == "KHAC" else None,
                "AC": _vnd(thue, 2),
            }
            if code not in ("", "KCT", "KKKNT"):
                row["AE"] = s.tk_thue
                row["AF"] = s.nhom_hhdv
            rows.append(row)

        lines = [ln for ln in inv.lines if ln.tchat != "4" and (ln.ten_hang or ln.thanh_tien)]
        sum_tt = 0.0
        sum_thue = 0.0
        if not lines:
            ts = inv.thue_theo_suat[0]["thue_suat"] if len(inv.thue_theo_suat) == 1 else ""
            ma = catalog.code_for("", "Hàng hóa, dịch vụ theo hóa đơn", "", ts, s)
            add_line(ma, "Hàng hóa, dịch vụ theo hóa đơn", "", None, None, inv.tong_tien_chua_thue, None, None, ts, inv.tong_tien_thue)
            notes.append("Không có dòng hàng chi tiết; ghi 1 dòng theo tổng tiền hóa đơn")
            sum_tt += inv.tong_tien_chua_thue or 0
            sum_thue += inv.tong_tien_thue or 0
            res.lines += 1
        for ln in lines:
            ma = catalog.code_for(ln.ma_hang, ln.ten_hang, ln.dvt, ln.thue_suat, s)
            sign = -1 if ln.tchat == "3" else 1  # chiết khấu thương mại ghi giảm
            tt = None if ln.thanh_tien is None else sign * ln.thanh_tien
            thue = None if ln.tien_thue is None else sign * ln.tien_thue
            if ln.tchat == "3":
                notes.append(f"Dòng {ln.stt} là chiết khấu thương mại, ghi âm")
            add_line(ma, ln.ten_hang, ln.dvt, ln.so_luong, ln.don_gia, tt, ln.tl_ck, ln.st_ck, ln.thue_suat, thue)
            sum_tt += tt or 0
            sum_thue += thue or 0
            res.lines += 1

        diff_tt = None if inv.tong_tien_chua_thue is None else round((inv.tong_tien_chua_thue or 0) - sum_tt, 2)
        diff_thue = None if inv.tong_tien_thue is None else round((inv.tong_tien_thue or 0) - sum_thue, 2)
        if (diff_tt and abs(diff_tt) >= 1) or (diff_thue and abs(diff_thue) >= 1):
            notes.append("Tổng dòng hàng lệch tổng hóa đơn, kiểm tra trước khi nhập")
        checks.append({
            "so_ct": so_ct, "hd": label, "ngay": inv.ngay_lap, "mst": inv.nb_mst,
            "tt_dong": _vnd(sum_tt, 2), "tt_hd": inv.tong_tien_chua_thue, "lech_tt": diff_tt,
            "thue_dong": _vnd(sum_thue, 2), "thue_hd": inv.tong_tien_thue, "lech_thue": diff_thue,
            "tong_tt": inv.tong_tien_tt, "co_xml": "Có" if has_xml else "Không",
            "ghi_chu": "; ".join(n for n in notes if n),
        })

        if has_xml:
            # Tên file: <MST người bán>_<ký hiệu>_<số>, lấy từ tên file gốc (purchase_<nguồn>_<MST>_<mẫu>_<ký hiệu>_<số>).
            parts = stem.split("_")
            nice = "_".join([parts[2], parts[4], "_".join(parts[5:])]) if len(parts) >= 6 else stem
            shutil.copyfile(path, xml_out / f"{nice}.xml")
            res.xml_files += 1
        res.vouchers += 1

    res.excel_path = str(misa_dir / TEMPLATE_OUT_NAME)
    _fill_template(Path(res.excel_path), rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "KiemTra"
    _write(ws, CHECK_COLUMNS, checks)
    _write(wb.create_sheet("DanhMuc_NhaCungCap"), SUPPLIER_COLUMNS, list(suppliers.values()))
    _write(wb.create_sheet("DanhMuc_VatTuHangHoa"), ITEM_COLUMNS, list(catalog.by_key.values()))
    _write(wb.create_sheet("BoQua"), [("Hóa đơn", "hd", False), ("MST người bán", "mst", False), ("Lý do không đưa vào", "ly_do", False), ("File", "file", False)], skipped)
    guide = wb.create_sheet("HuongDan")
    for line in HUONG_DAN.strip().splitlines():
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    res.catalog_path = str(misa_dir / "MISA_DanhMuc_va_KiemTra.xlsx")
    wb.save(res.catalog_path)
    for old in (misa_dir / "MISA_NhapKhau_MuaHang.xlsx", misa_dir / "LOI_XUAT_MISA.txt"):  # file cũ
        if old.exists():
            old.unlink()

    res.skipped = len(skipped)
    res.suppliers = len(suppliers)
    res.items = len(catalog.by_key)
    res.warnings = [c["hd"] + ": " + c["ghi_chu"] for c in checks if "lệch" in c["ghi_chu"]]
    (misa_dir / "HUONG_DAN_NHAP_MISA.txt").write_text(HUONG_DAN.strip() + "\n", encoding="utf-8")
    log.info(
        "[MISA] %d chứng từ (%d dòng), %d XML, %d hóa đơn không có XML, %d bỏ qua; %d NCC, %d mã hàng -> %s",
        res.vouchers, res.lines, res.xml_files, res.no_xml, res.skipped, res.suppliers, res.items, res.excel_path,
    )
    return res


CHECK_COLUMNS: list[tuple[str, str, bool]] = [
    ("Số chứng từ", "so_ct", False), ("Hóa đơn", "hd", False), ("Ngày", "ngay", False), ("MST người bán", "mst", False),
    ("Tổng thành tiền các dòng", "tt_dong", False), ("Tiền hàng trên HĐ", "tt_hd", False), ("Lệch tiền hàng", "lech_tt", False),
    ("Tổng thuế các dòng", "thue_dong", False), ("Tiền thuế trên HĐ", "thue_hd", False), ("Lệch tiền thuế", "lech_thue", False),
    ("Tổng thanh toán HĐ", "tong_tt", False), ("Có XML gốc", "co_xml", False), ("Ghi chú", "ghi_chu", False),
]


HUONG_DAN = r"""
HƯỚNG DẪN NHẬP HÓA ĐƠN MUA VÀO VÀO MISA SME
===========================================

Thư mục misa gồm:
- Mua_hang_khong_qua_kho_VND.xlsx : ĐÚNG MẪU nhập khẩu "Chứng từ mua hàng không qua kho" (VND) của MISA.
                                    Mỗi dòng hàng một dòng; các dòng cùng "Số chứng từ" là một chứng từ.
- MISA_DanhMuc_va_KiemTra.xlsx    : KiemTra (đối chiếu tổng dòng hàng với tổng hóa đơn, ghi chú cần xem lại),
                                    DanhMuc_NhaCungCap, DanhMuc_VatTuHangHoa, BoQua (hóa đơn bị hủy/bị thay thế).
- xml_mua_vao\                    : XML gốc, dùng nếu muốn MISA tự đọc hóa đơn điện tử (không bắt buộc).

TRƯỚC KHI NHẬP
1. Mở MISA_DanhMuc_va_KiemTra.xlsx, sheet KiemTra: xem các dòng có ghi chú "lệch", "kiểm tra lại".
2. Mở Mua_hang_khong_qua_kho_VND.xlsx, kiểm tra các cột tài khoản:
   - "TK chi phí (*)"         mặc định 642 (TT133). Dùng TT200 thì thường là 6422; hàng mua về bán thì dùng mẫu qua kho.
   - "TK công nợ/TK tiền (*)" mặc định 331 (Phương thức thanh toán = 0: chưa thanh toán).
   - "TK thuế GTGT"           mặc định 1331; "Nhóm HHDV mua vào" mặc định 1.
   Đơn vị kinh doanh ngành không chịu thuế GTGT (vd giáo dục) cần xem lại thuế đầu vào có được khấu trừ không.
   Có thể đổi mặc định trong file .env (MISA_TK_CHI_PHI, MISA_TK_CONG_NO, MISA_TK_THUE, MISA_NHOM_HHDV,
   MISA_PHUONG_THUC_TT) rồi bấm lại "Xuất sang MISA".

BƯỚC 1 - NHẬP DANH MỤC (chỉ cần với nhà cung cấp / mặt hàng chưa có trong MISA)
   MISA: Tệp > Nhập khẩu từ Excel > Danh mục Nhà cung cấp, chọn MISA_DanhMuc_va_KiemTra.xlsx, sheet DanhMuc_NhaCungCap.
   Làm tương tự với Danh mục Vật tư hàng hóa, sheet DanhMuc_VatTuHangHoa. Ở bước "Ghép dữ liệu" kiểm tra cột ghép đúng.
   Mã nhà cung cấp = MST người bán. Nếu MISA đã có NCC với mã khác, sửa cột "Mã nhà cung cấp" trong file chứng từ.

BƯỚC 2 - NHẬP CHỨNG TỪ
   MISA: Mua hàng > tab Mua hàng hóa, dịch vụ > mũi tên cạnh "Thêm chứng từ mua hàng" > Nhập từ excel
   (hoặc Tệp > Nhập khẩu từ Excel > Chứng từ mua hàng không qua kho).
   Chọn file Mua_hang_khong_qua_kho_VND.xlsx. Vì file đúng mẫu MISA nên các cột tự ghép.
   Bấm Thực hiện, xem kết quả kiểm tra, sửa dòng MISA báo lỗi rồi nhập lại.
   Nếu MISA chỉ nhận .xls: mở file bằng Excel > Save As > Excel 97-2003 Workbook (.xls).

LƯU Ý
- Số chứng từ tạo dạng MH + ký hiệu + số hóa đơn (tối đa 20 ký tự) để không trùng.
- Cột ĐVT phải có trong danh mục Đơn vị tính của MISA; ĐVT lạ cần thêm vào MISA trước.
- Chiết khấu thương mại ghi âm; hàng khuyến mại giữ số liệu trên hóa đơn.
- Không vừa nhập Excel vừa cho MISA đọc XML cho cùng một hóa đơn (sẽ trùng).
"""
