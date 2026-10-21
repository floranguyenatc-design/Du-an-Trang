"""Tạo file nhập khẩu MISA SME cho hóa đơn BÁN RA đã kéo về.

Kết quả trong thư mục ``misa``:

- ``Ban_hang_VND.xlsx``: điền mẫu nhập khẩu "Chứng từ bán hàng" (Bán hàng VND) của MISA, ghi theo
  tên cột. Mỗi dòng hàng một dòng, mỗi hóa đơn một chứng từ (số chứng từ BH + số hóa đơn).
- ``Danh_muc_khach_hang.xlsx``: điền mẫu "Danh mục KH" của MISA; nhập vào MISA trước file bán hàng.
  Mã khách hàng: MST nếu khách là đơn vị có MST, ngược lại là tên không dấu viết liền (vd NGUYENVANAN).
- ``MISA_BanHang_KiemTra.xlsx``: đối chiếu tổng dòng hàng với tổng hóa đơn, ghi chú cần xem lại,
  danh mục mã hàng, hóa đơn không đưa vào (bị hủy, bị thay thế, chưa được cấp mã...).
- ``MaHang_MISA.xlsx`` (dùng chung với chiều mua vào): sửa mã hàng một lần, lần sau tự dùng lại.
"""

from __future__ import annotations

import logging
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .misa import (
    MAP_FILE, TEMPLATES_DIR, ItemCatalog, MaHangMap, MisaSettings, _collect, _cut, _load_raw_index,
    _norm_name, _vnd, _write, huong_dan, skip_reason, template_info,
)
from . import misa as _misa
from .thuesuat import misa_vat_rate

log = logging.getLogger("hddt")

TEMPLATE_BAN = TEMPLATES_DIR / "misa_ban_hang_vnd.xlsx"
TEMPLATE_DM_KH = TEMPLATES_DIR / "misa_danh_muc_khach_hang.xlsx"
OUT_BAN = "Ban_hang_VND.xlsx"
OUT_DM_KH = "Danh_muc_khach_hang.xlsx"
OUT_CHECK = "MISA_BanHang_KiemTra.xlsx"

# Phương thức thanh toán của mẫu bán hàng: 0 chưa thu tiền, 1 thu tiền mặt, 2 thu chuyển khoản.
PAY_LABEL_BAN = {"0": "chưa thu tiền", "1": "thu tiền mặt", "2": "thu chuyển khoản"}

DATE_HEADERS = {"Ngày hạch toán (*)", "Ngày chứng từ (*)", "Ngày hóa đơn", "Ngày cấp"}
MAX_LEN = {"Số chứng từ (*)": 20, "Số hóa đơn": 25, "Tên khách hàng": 128, "Tên khách hàng (*)": 128, "Địa chỉ": 255,
           "Diễn giải": 255, "Tên hàng": 255, "Mã số thuế": 50, "TK Tiền/Chi phí/Nợ (*)": 20, "TK Doanh thu/Có (*)": 20}
NUMERIC_TYPES = (int, float)

CHECK_COLUMNS = [
    ("Số chứng từ", "so_ct", False), ("Hóa đơn", "hd", False), ("Ngày", "ngay", False), ("Mã khách hàng", "ma_kh", False),
    ("Thu tiền / TK (cột C, Y)", "tt", False),
    ("Tổng thành tiền các dòng", "tt_dong", False), ("Tiền hàng trên HĐ", "tt_hd", False), ("Lệch tiền hàng", "lech_tt", False),
    ("Tổng thuế các dòng", "thue_dong", False), ("Tiền thuế trên HĐ", "thue_hd", False), ("Lệch tiền thuế", "lech_thue", False),
    ("Có XML gốc", "co_xml", False), ("Ghi chú", "ghi_chu", False),
]
ITEM_COLUMNS = [("Mã hàng", "ma", True), ("Tên hàng", "ten", True), ("ĐVT", "dvt", False), ("Thuế suất", "thue_suat", False)]


@dataclass
class BanResult:
    folder: str = ""
    excel_path: str = ""
    dm_kh_path: str = ""
    check_path: str = ""
    vouchers: int = 0
    lines: int = 0
    customers: int = 0
    items: int = 0
    skipped: int = 0
    no_xml: int = 0
    warnings: list[str] = field(default_factory=list)


def customer_code(name: str) -> str:
    """Tên không dấu viết liền, in hoa: 'Nguyễn Ngọc Bảo Hân' -> 'NGUYENNGOCBAOHAN'."""
    t = unicodedata.normalize("NFD", name or "")
    t = "".join(ch for ch in t if unicodedata.category(ch) != "Mn").replace("đ", "d").replace("Đ", "D")
    return re.sub(r"[^A-Za-z0-9]", "", t).upper()[:25] or "KH"


def _fill(template: Path, out: Path, rows: list[dict[str, Any]]) -> list[str]:
    wb = load_workbook(template)
    ws = wb.active
    head = {str(c.value).strip(): c.column_letter for c in ws[1] if c.value}
    lost: dict[str, int] = {}
    for i, row in enumerate(rows, start=2):
        for h, v in row.items():
            if v in (None, ""):
                continue
            col = head.get(h)
            if col is None:
                lost[h] = lost.get(h, 0) + 1
                continue
            if h in MAX_LEN and isinstance(v, str):
                v = _cut(v, MAX_LEN[h])
            cell = ws[f"{col}{i}"]
            cell.value = v
            if h in DATE_HEADERS:
                cell.number_format = "dd/mm/yyyy"
            elif isinstance(v, str):
                cell.number_format = "@"
            elif isinstance(v, NUMERIC_TYPES) and not isinstance(v, bool) and abs(v) >= 1000:
                cell.number_format = "#,##0.####"
    wb.save(out)
    return [f"Mẫu MISA không có cột '{h}': bỏ qua {n} ô" for h, n in lost.items()]


def export_misa_ban(output_dir: str | Path, settings: MisaSettings | None = None) -> BanResult:
    s = settings or MisaSettings()
    out_dir = Path(output_dir)
    misa_dir = out_dir / "misa"
    misa_dir.mkdir(parents=True, exist_ok=True)
    res = BanResult(folder=str(misa_dir))

    raw_index = _load_raw_index(out_dir, "sold")
    invoices = _collect(out_dir, raw_index, "sold")
    if not invoices:
        raise FileNotFoundError(f"Không thấy hóa đơn bán ra trong {out_dir / 'xml' / 'sold'}. Hãy kéo hóa đơn bán ra (có tải XML) trước.")
    invoices.sort(key=lambda t: (t[2].ngay_lap or datetime.min, t[2].khhdon, int(t[2].shdon) if str(t[2].shdon).isdigit() else 0))

    _, allowed_vat = template_info(TEMPLATE_BAN)
    pay = str(s.ban_thanh_toan).strip() if str(s.ban_thanh_toan).strip() in PAY_LABEL_BAN else "1"
    tk_no = {"0": s.tk_phai_thu, "1": s.tk_tien_mat, "2": s.tk_ngan_hang}[pay]

    mahang = MaHangMap(misa_dir / MAP_FILE)
    catalog = ItemCatalog()
    customers: dict[str, dict[str, Any]] = {}   # norm tên/MST -> thông tin
    used_codes: dict[str, str] = {}             # mã -> khóa khách hàng
    rows: list[dict[str, Any]] = []
    checks: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    used_so_ct: set[str] = set()

    for stem, path, inv, has_xml in invoices:
        raw = raw_index.get(stem, {})
        ten_kh = (inv.nm_ten or inv.nm_hvtn or str(raw.get("nmten") or raw.get("nmtnmua") or "")).strip()
        ten_kh = re.sub(r"\s+", " ", ten_kh)
        so = str(raw.get("shdon") or inv.shdon or "").strip()
        kh = inv.khhdon or str(raw.get("khhdon") or "")
        label = f"{inv.khmshdon}{kh} số {so} - {ten_kh}"
        reason = skip_reason(raw)
        if reason:
            skipped.append({"hd": label, "mst": inv.nm_mst, "ly_do": reason, "file": path.name})
            continue

        # ---- khách hàng
        mst = (inv.nm_mst or "").strip()
        ckey = f"mst:{mst}" if mst else f"ten:{_norm_name(ten_kh)}"
        cus = customers.get(ckey)
        if cus is None:
            base = mst if mst else customer_code(ten_kh)
            code, n = base, 1
            while code in used_codes and used_codes[code] != ckey:
                n += 1
                code = f"{base[:23]}{n:02d}"
            used_codes[code] = ckey
            cus = customers[ckey] = {"ma": code, "ten": ten_kh, "mst": mst, "dia_chi": None, "email": None, "ca_nhan": 0 if mst else 1, "so_hd": []}
        if inv.nm_dchi:
            cus["dia_chi"] = re.sub(r"\s+", " ", inv.nm_dchi.strip())  # hóa đơn sau ghi đè: địa chỉ mới nhất
        if inv.nm_email:
            cus["email"] = inv.nm_email
        cus["so_hd"].append(so)

        # ---- số chứng từ
        so_ct = f"{s.ban_so_ct_prefix}{int(so):05d}" if so.isdigit() else f"{s.ban_so_ct_prefix}{so}"
        if so_ct in used_so_ct:
            so_ct = f"{s.ban_so_ct_prefix}{re.sub(r'[^A-Za-z0-9]', '', kh)}{so}"[:20]
        n = 1
        while so_ct in used_so_ct:
            n += 1
            so_ct = f"{so_ct[:17]}_{n}"
        used_so_ct.add(so_ct)

        notes: list[str] = []
        if not has_xml:
            notes.append("Không có XML gốc; số liệu lấy từ chi tiết trên cổng thuế")
            res.no_xml += 1
        if _int_or0(raw.get("tthai")) in (2, 3, 5):
            notes.append("Hóa đơn thay thế/điều chỉnh/bị điều chỉnh - kiểm tra lại trước khi hạch toán")

        header = {
            "Hình thức bán hàng": 0, "Phương thức thanh toán": int(pay), "Kiêm phiếu xuất kho": 0, "Lập kèm hóa đơn": 1,
            "Ngày hạch toán (*)": inv.ngay_lap, "Ngày chứng từ (*)": inv.ngay_lap, "Số chứng từ (*)": so_ct,
            "Số hóa đơn": so, "Ngày hóa đơn": inv.ngay_lap,
            "Mã khách hàng": cus["ma"], "Tên khách hàng": ten_kh, "Địa chỉ": inv.nm_dchi or None, "Mã số thuế": mst or None,
            "Diễn giải": f"Bán hàng cho {ten_kh} theo HĐ {kh} số {so}",
            "TK Tiền/Chi phí/Nợ (*)": tk_no, "TK Doanh thu/Có (*)": s.tk_doanh_thu,
        }

        def add_line(ma_goc, ten_hang, dvt, sl, dg, tt, tl_ck, st_ck, ts, thue, khuyen_mai=False):
            code, _khac = misa_vat_rate(ts)
            if code and allowed_vat is not None and code not in allowed_vat:
                msg = (f"Thuế suất {code} không có trong ghi chú của mẫu MISA (chỉ có "
                       f"{', '.join(sorted(allowed_vat, key=lambda x: (not x.isdigit(), int(x) if x.isdigit() else 0, x)))}); kiểm tra khi nhập")
                if msg not in notes:
                    notes.append(msg)
            ma = catalog.code_for(ma_goc, ten_hang, dvt, ts, s, override=mahang.get("Bán ra", ten_hang, dvt))
            mahang.note("Bán ra", ten_hang, dvt, ma)
            row = {**header, "Mã hàng (*)": ma, "Tên hàng": ten_hang, "ĐVT": dvt or None,
                   "Số lượng": _vnd(sl, 4), "Đơn giá": _vnd(dg, 4), "Thành tiền": _vnd(tt, 2),
                   "Tỷ lệ CK (%)": _vnd(tl_ck, 4) if tl_ck else None, "Tiền chiết khấu": _vnd(st_ck, 2) if st_ck else None,
                   "% thuế GTGT": code or None, "Tiền thuế GTGT": _vnd(thue, 2) if thue is not None else (0 if code else None),
                   "Hàng khuyến mại": 1 if khuyen_mai else None}
            if code and code not in ("KCT", "KKKNT"):
                row["TK thuế GTGT"] = s.tk_thue_ra
            rows.append(row)
            res.lines += 1

        lines = [ln for ln in inv.lines if ln.tchat != "4" and (ln.ten_hang or ln.thanh_tien)]
        sum_tt = sum_thue = 0.0
        if not lines:
            ts = inv.thue_theo_suat[0]["thue_suat"] if len(inv.thue_theo_suat) == 1 else ""
            add_line("", "Hàng hóa, dịch vụ theo hóa đơn", "", None, None, inv.tong_tien_chua_thue, None, None, ts, inv.tong_tien_thue)
            notes.append("Không có dòng hàng chi tiết; ghi 1 dòng theo tổng tiền hóa đơn")
            sum_tt += inv.tong_tien_chua_thue or 0
            sum_thue += inv.tong_tien_thue or 0
        for ln in lines:
            sign = -1 if ln.tchat == "3" else 1  # chiết khấu thương mại ghi giảm
            tt = None if ln.thanh_tien is None else sign * ln.thanh_tien
            thue = None if ln.tien_thue is None else sign * ln.tien_thue
            if ln.tchat == "3":
                notes.append(f"Dòng {ln.stt} là chiết khấu thương mại, ghi âm")
            add_line(ln.ma_hang, ln.ten_hang, ln.dvt, ln.so_luong, ln.don_gia, tt, ln.tl_ck, ln.st_ck, ln.thue_suat, thue,
                     khuyen_mai=(ln.tchat == "2"))
            sum_tt += tt or 0
            sum_thue += thue or 0

        lech_tt = None if inv.tong_tien_chua_thue is None else round((inv.tong_tien_chua_thue or 0) - sum_tt, 2)
        lech_thue = None if inv.tong_tien_thue is None else round((inv.tong_tien_thue or 0) - sum_thue, 2)
        if (lech_tt and abs(lech_tt) >= 1) or (lech_thue and abs(lech_thue) >= 1):
            notes.append("Tổng dòng hàng lệch tổng hóa đơn, kiểm tra trước khi nhập")
        checks.append({
            "so_ct": so_ct, "hd": label, "ngay": inv.ngay_lap, "ma_kh": cus["ma"],
            "tt": f"{PAY_LABEL_BAN[pay]}, Nợ {tk_no} / Có {s.tk_doanh_thu}",
            "tt_dong": _vnd(sum_tt, 2), "tt_hd": inv.tong_tien_chua_thue, "lech_tt": lech_tt,
            "thue_dong": _vnd(sum_thue, 2), "thue_hd": inv.tong_tien_thue, "lech_thue": lech_thue,
            "co_xml": "Có" if has_xml else "Không", "ghi_chu": "; ".join(notes),
        })
        res.vouchers += 1

    # ---- ghi file
    res.excel_path = str(misa_dir / OUT_BAN)
    res.warnings += _fill(TEMPLATE_BAN, Path(res.excel_path), rows)

    kh_rows = [{"Là tổ chức/cá nhân": c["ca_nhan"], "Mã khách hàng (*)": c["ma"], "Tên khách hàng (*)": c["ten"],
                "Địa chỉ": c["dia_chi"], "Mã số thuế": c["mst"] or None, "Email": c["email"]}
               for c in sorted(customers.values(), key=lambda c: c["ma"])]
    res.dm_kh_path = str(misa_dir / OUT_DM_KH)
    res.warnings += _fill(TEMPLATE_DM_KH, Path(res.dm_kh_path), kh_rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "KiemTra"
    _write(ws, CHECK_COLUMNS, checks)
    _write(wb.create_sheet("DanhMuc_HangHoa"), ITEM_COLUMNS, list(catalog.by_key.values()))
    _write(wb.create_sheet("BoQua"), [("Hóa đơn", "hd", False), ("MST người mua", "mst", False), ("Lý do không đưa vào", "ly_do", False), ("File", "file", False)], skipped)
    guide = wb.create_sheet("HuongDan")
    for line in huong_dan().splitlines():
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    res.check_path = str(misa_dir / OUT_CHECK)
    wb.save(res.check_path)
    w = mahang.save()
    if w:
        res.warnings.append(w)
    (misa_dir / "HUONG_DAN_NHAP_MISA.txt").write_text(huong_dan() + "\n", encoding="utf-8")

    res.customers = len(customers)
    res.items = len(catalog.by_key)
    res.skipped = len(skipped)
    res.warnings += [c["hd"] + ": " + c["ghi_chu"] for c in checks if "lệch" in c["ghi_chu"] or "mẫu MISA" in c["ghi_chu"]]
    log.info("[MISA bán ra] %d chứng từ (%d dòng), %d khách hàng, %d mã hàng, %d bỏ qua -> %s",
             res.vouchers, res.lines, res.customers, res.items, res.skipped, res.excel_path)
    return res


def _int_or0(v: Any) -> int:
    try:
        return int(v) if v not in (None, "") else 0
    except (TypeError, ValueError):
        return 0


@dataclass
class MisaAll:
    folder: str = ""
    mua: Any = None          # MisaResult
    ban: BanResult | None = None
    errors: dict[str, str] = field(default_factory=dict)  # chiều -> lỗi

    @property
    def vouchers(self) -> int:
        return (self.mua.vouchers if self.mua else 0) + (self.ban.vouchers if self.ban else 0)

    def summary(self) -> str:
        parts = []
        if self.mua:
            parts.append(f"mua vào {self.mua.vouchers} chứng từ (Mua_hang_khong_qua_kho_VND.xlsx)")
        if self.ban:
            parts.append(f"bán ra {self.ban.vouchers} chứng từ, {self.ban.customers} khách hàng (Ban_hang_VND.xlsx, Danh_muc_khach_hang.xlsx)")
        return "; ".join(parts)


def export_misa_all(output_dir: str | Path, settings: MisaSettings | None = None,
                    directions: tuple[str, ...] = ("purchase", "sold"), only_existing: bool = True) -> MisaAll:
    """Xuất MISA cho các chiều có dữ liệu. only_existing: bỏ qua chiều chưa có thư mục xml (không báo lỗi)."""
    out_dir = Path(output_dir)
    res = MisaAll(folder=str(out_dir / "misa"))
    todo = [d for d in directions if not only_existing or (out_dir / "xml" / d).is_dir()]
    if not todo:
        raise FileNotFoundError(
            f"Không thấy hóa đơn đã kéo trong {out_dir / 'xml'}. Hãy kéo hóa đơn (có tải XML) trước rồi mới xuất sang MISA.")
    for d in todo:
        try:
            if d == "purchase":
                res.mua = _misa.export_misa(out_dir, settings)
            else:
                res.ban = export_misa_ban(out_dir, settings)
        except Exception as exc:  # noqa: BLE001
            res.errors[d] = str(exc) or type(exc).__name__
            log.error("[MISA %s] %s", "mua vào" if d == "purchase" else "bán ra", res.errors[d])
            log.debug("Chi tiết lỗi MISA", exc_info=True)
    return res
