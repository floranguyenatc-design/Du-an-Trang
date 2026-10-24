import json
import shutil
from pathlib import Path

import xlrd

from hddt.bangke import Item, export_bang_ke, write_bang_ke

FIX = Path(__file__).parent / "fixtures" / "hoadon_mau.xml"


def _sections(path):
    sh = xlrd.open_workbook(str(path), formatting_info=True).sheet_by_name("Bang ke")
    rows = [[sh.cell_value(r, c) for c in range(1, 7)] for r in range(sh.nrows)]
    i2 = next(i for i, r in enumerate(rows) if str(r[0]).startswith("II. "))
    data = lambda a, b: [r for r in rows[a:b] if isinstance(r[0], float)]  # noqa: E731
    return data(0, i2), data(i2, len(rows)), sh


def _setup(tmp_path: Path) -> Path:
    out = tmp_path / "out"
    pur, sold = out / "xml" / "purchase", out / "xml" / "sold"
    pur.mkdir(parents=True)
    sold.mkdir(parents=True)
    shutil.copy(FIX, pur / "purchase_query_0312345678_1_C25TAA_125.xml")
    shutil.copy(FIX, pur / "purchase_query_0312345678_1_C25TAA_130.xml")  # bị hủy
    (out / "danh_sach_muavao_20251201_20251231.json").write_text(json.dumps([
        {"nbmst": "0312345678", "khmshdon": "1", "khhdon": "C25TAA", "shdon": "125", "tthai": 1, "ttxly": 5},
        {"nbmst": "0312345678", "khmshdon": "1", "khhdon": "C25TAA", "shdon": "130", "tthai": 6, "ttxly": 5},
    ]), encoding="utf-8")
    sold_items = []
    for so in ("12", "13"):
        (sold / f"sold_query_0109999999_1_C25TBB_{so}.xml").write_text(
            FIX.read_text(encoding="utf-8").replace("<SHDon>125</SHDon>", f"<SHDon>{so}</SHDon>"), encoding="utf-8")
        sold_items.append({"nbmst": "0109999999", "khmshdon": "1", "khhdon": "C25TBB", "shdon": so, "tthai": 1, "ttxly": 5})
    (out / "danh_sach_banra_20251201_20251231.json").write_text(json.dumps(sold_items), encoding="utf-8")
    return out


def test_export_bang_ke_quarterly(tmp_path):
    out = _setup(tmp_path)
    res = export_bang_ke(out)
    assert [Path(f).name for f in res.files] == ["BangKe01_GiamThue_Quy4_2025.xls"]
    mua, ban, sh = _sections(res.files[0])
    # Phần I: chỉ dòng 8% của hóa đơn hợp lệ (bỏ hóa đơn bị hủy, bỏ dòng 10%)
    assert mua == [[1.0, "Văn phòng phẩm", 500000.0, 40000.0, "", ""]]
    # Phần II: 2 hóa đơn cùng tên hàng được cộng gộp
    assert ban == [[1.0, "Văn phòng phẩm", 1000000.0, 10.0, 8.0, 20000.0]]
    assert [s.name for s in xlrd.open_workbook(res.files[0]).sheets()] == ["Bang ke", "HDSD"]
    assert res.periods[0]["giam"] == 20000

    from openpyxl import load_workbook

    wb = load_workbook(res.detail_path)
    assert wb.sheetnames == ["TongHop", "ChiTiet", "KhongDuaVao"]
    assert wb["ChiTiet"].max_row == 4  # 1 mua vào + 2 bán ra
    assert "hủy" in str(wb["KhongDuaVao"]["E2"].value)


def test_export_bang_ke_monthly(tmp_path):
    res = export_bang_ke(_setup(tmp_path), "thang")
    assert [Path(f).name for f in res.files] == ["BangKe01_GiamThue_Thang12_2025.xls"]


def test_many_rows_shift_section_ii(tmp_path):
    p = tmp_path / "bk.xls"
    write_bang_ke(p, [Item(f"Hàng mua {i}", 1000 * i, 80 * i) for i in range(1, 8)], [Item(f"Hàng bán {i}", 100000, 8000) for i in range(1, 7)])
    mua, ban, sh = _sections(p)
    assert [r[0] for r in mua] == [float(i) for i in range(1, 8)] and mua[6][1:4] == ["Hàng mua 7", 7000.0, 560.0]
    assert len(ban) == 6 and ban[5][1:] == ["Hàng bán 6", 100000.0, 10.0, 8.0, 2000.0]
    # tiêu đề phần II (ô gộp 3 dòng) dời xuống theo số dòng phần I
    head = [c for c in sh.merged_cells if c[2] == 6]
    assert head and head[0][0] > 29 + 7


def test_no_8_percent_no_file(tmp_path):
    out = _setup(tmp_path)
    for f in (out / "xml").rglob("*.xml"):
        f.write_text(f.read_text(encoding="utf-8").replace("<TSuat>8%</TSuat>", "<TSuat>10%</TSuat>"), encoding="utf-8")
    res = export_bang_ke(out)
    assert res.files == [] and not (out / "bang_ke_giam_thue").exists()


def test_cli_bang_ke(tmp_path, capsys):
    from hddt import cli

    out = _setup(tmp_path)
    assert cli.main(["--env", "x", "bang-ke", "--thu-muc", str(out)]) == 0
    assert "Quý 4/2025" in capsys.readouterr().out
