"""Chạy trọn luồng `pull` và `parse-xml` với server giả lập."""

from pathlib import Path
from unittest import mock

import requests
from openpyxl import load_workbook

from hddt import cli
from tests.test_client import FakeGdt

FIX = Path(__file__).parent / "fixtures" / "hoadon_mau.xml"
RealSession = requests.Session


def _run(argv, fake, monkeypatch):
    session = mock.Mock(spec=RealSession)
    session.request.side_effect = fake.request
    session.proxies = {}
    monkeypatch.setattr(requests, "Session", lambda: session)
    monkeypatch.setenv("GDT_USERNAME", "0109876543")
    monkeypatch.setenv("GDT_PASSWORD", "secret")
    monkeypatch.setenv("REQUEST_DELAY_MS", "0")
    monkeypatch.setenv("PAGE_SIZE", "2")
    return cli.main(argv)


def test_pull_end_to_end(tmp_path, monkeypatch):
    fake = FakeGdt()
    out = tmp_path / "out"
    code = _run(["--env", str(tmp_path / "no.env"), "pull", "--thang", "12/2025", "--chieu", "mua", "--thu-muc", str(out), "--luong", "2"], fake, monkeypatch)
    assert code == 2  # có 1 hóa đơn không tải được XML (HTTP 500) -> mã thoát 2
    xlsx = out / "HoaDon_20251201_20251231.xlsx"
    assert xlsx.is_file()
    wb = load_workbook(xlsx)
    ws = wb["HoaDon"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:5] == ("Chiều", "Nguồn", "Ký hiệu mẫu số", "Ký hiệu hóa đơn", "Số hóa đơn")
    assert len(rows) == 5  # 1 header + 4 hóa đơn
    by_no = {int(r[4]): r for r in rows[1:]}
    assert by_no[125][18] == "Hóa đơn mới" and by_no[125][19] == "Đã cấp mã hóa đơn"
    assert by_no[9][1] == "Máy tính tiền"
    assert "HTTP 500" in by_no[126][-1]
    assert (out / "xml" / "purchase" / "purchase_query_0312345678_1_C25TAA_125.xml").is_file()
    assert (out / "xml" / "purchase" / "purchase_query_0312345678_1_C25TAA_125.html").is_file()
    err_rows = list(wb["Loi"].iter_rows(values_only=True))
    assert len(err_rows) == 2 and err_rows[1][5] == "126"
    assert (out / "danh_sach_20251201_20251231.json").is_file()


def test_pull_reuses_existing_xml(tmp_path, monkeypatch):
    fake = FakeGdt()
    out = tmp_path / "out"
    _run(["--env", "x", "pull", "--thang", "12/2025", "--chieu", "mua", "--thu-muc", str(out), "--khong-mtt"], fake, monkeypatch)
    first = sum(1 for c in fake.calls if "export-xml" in c[1])
    _run(["--env", "x", "pull", "--thang", "12/2025", "--chieu", "mua", "--thu-muc", str(out), "--khong-mtt"], fake, monkeypatch)
    second = sum(1 for c in fake.calls if "export-xml" in c[1]) - first
    assert first == 3 and second == 1  # lần 2 chỉ tải lại hóa đơn 126 bị lỗi


def test_parse_xml_offline(tmp_path, monkeypatch):
    folder = tmp_path / "xml" / "purchase"
    folder.mkdir(parents=True)
    (folder / "a.xml").write_bytes(FIX.read_bytes())
    out = tmp_path / "local.xlsx"
    code = cli.main(["--env", "x", "parse-xml", str(tmp_path / "xml"), "--excel", str(out)])
    assert code == 0
    wb = load_workbook(out)
    assert wb["HoaDon"].max_row == 2
    assert wb["ChiTiet"].max_row == 4
    detail = list(wb["ChiTiet"].iter_rows(min_row=2, values_only=True))
    assert detail[0][10] == "Dịch vụ tư vấn" and detail[0][18] == 200000
