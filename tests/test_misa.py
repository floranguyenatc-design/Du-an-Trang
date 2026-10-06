import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from hddt.misa import MisaSettings, export_misa

FIX = Path(__file__).parent / "fixtures" / "hoadon_mau.xml"


def _setup(tmp_path: Path) -> Path:
    out = tmp_path / "out"
    pur = out / "xml" / "purchase"
    pur.mkdir(parents=True)
    items = [
        # có XML, hợp lệ
        {"nbmst": "0312345678", "khmshdon": "1", "khhdon": "C25TAA", "shdon": "125", "tthai": 1, "ttxly": 5},
        # không có XML, chi tiết từ GDT
        {"nbmst": "0100109106", "khmshdon": "1", "khhdon": "K24DAA", "shdon": "6928336", "tthai": 1, "ttxly": 6},
        # bị hủy -> bỏ qua
        {"nbmst": "0312345678", "khmshdon": "1", "khhdon": "C25TAA", "shdon": "130", "tthai": 6, "ttxly": 5},
        # gói có 2 XML -> lưu _1, _2
        {"nbmst": "0312345678", "khmshdon": "1", "khhdon": "C25TAA", "shdon": "140", "tthai": 3, "ttxly": 5},
    ]
    (out / "danh_sach_20251201_20251231.json").write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")
    shutil.copy(FIX, pur / "purchase_query_0312345678_1_C25TAA_125.xml")
    shutil.copy(FIX, pur / "purchase_query_0312345678_1_C25TAA_130.xml")
    xml140 = FIX.read_text(encoding="utf-8").replace("<SHDon>125</SHDon>", "<SHDon>140</SHDon>").replace(
        "<TChat>1</TChat><STT>2</STT>", "<TChat>3</TChat><STT>2</STT>"
    )
    (pur / "purchase_query_0312345678_1_C25TAA_140_1.xml").write_text(xml140, encoding="utf-8")
    (pur / "purchase_query_0312345678_1_C25TAA_140_2.xml").write_text(xml140, encoding="utf-8")
    detail = {
        "nbmst": "0100109106", "nbten": "Tập đoàn Công nghiệp - Viễn thông Quân đội", "khmshdon": 1, "khhdon": "K24DAA",
        "shdon": 6928336, "tdlap": "2024-04-02T00:00:00", "tgtcthue": 435729, "tgtthue": 43573, "tgtttbso": 479302,
        "hdhhdvu": [{"stt": 1, "ten": "Cước dịch vụ viễn thông", "dvtinh": "Tháng", "sluong": 1, "dgia": 435729,
                     "thtien": 435729, "tsuat": 0.1, "mhhdvu": "SP.0000"}],
    }
    (pur / "purchase_query_0100109106_1_K24DAA_6928336.json").write_text(json.dumps(detail, ensure_ascii=False), encoding="utf-8")
    return out


def test_export_misa(tmp_path):
    out = _setup(tmp_path)
    res = export_misa(out, MisaSettings(tk_chi_phi="6422"))
    assert res.vouchers == 3 and res.skipped == 1 and res.no_xml == 1 and res.xml_files == 2

    xmls = sorted(p.name for p in Path(res.xml_dir).glob("*.xml"))
    assert xmls == ["0312345678_C25TAA_125.xml", "0312345678_C25TAA_140.xml"]

    wb = load_workbook(res.excel_path)
    assert wb.sheetnames == ["ChungTuMuaHang", "DanhMuc_NhaCungCap", "DanhMuc_VatTuHangHoa", "BoQua", "HuongDan"]
    ws = wb["ChungTuMuaHang"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    # 125: 2 dòng hàng (bỏ dòng ghi chú TChat=4); 140: 2 dòng; Viettel: 1 dòng
    assert len(rows) == 5
    so_ct = {r["Số hóa đơn"]: r["Số chứng từ"] for r in rows}
    assert len(set(so_ct.values())) == 3
    assert all(r["TK kho/TK chi phí"] == "6422" and r["TK công nợ"] == "331" and r["TK thuế GTGT"] == "1331" for r in rows)

    viettel = [r for r in rows if r["Mã nhà cung cấp"] == "0100109106"][0]
    assert viettel["Tên hàng"] == "Cước dịch vụ viễn thông" and viettel["Đơn giá"] == 435729 and viettel["% thuế GTGT"] == "10%"
    assert "Không có XML gốc" in viettel["Ghi chú (tool)"]

    ck = [r for r in rows if r["Số hóa đơn"] == "140" and r["Tên hàng"] == "Văn phòng phẩm"][0]
    assert ck["Thành tiền"] == -500000 and ck["Tiền thuế GTGT"] == -40000 and "Chiết khấu" in ck["Ghi chú (tool)"]
    assert "kiểm tra lại" in ck["Ghi chú (tool)"]  # hóa đơn điều chỉnh

    ncc = list(wb["DanhMuc_NhaCungCap"].iter_rows(min_row=2, values_only=True))
    assert sorted(r[0] for r in ncc) == ["0100109106", "0312345678"]
    items = {r[1]: r for r in wb["DanhMuc_VatTuHangHoa"].iter_rows(min_row=2, values_only=True)}
    assert items["Cước dịch vụ viễn thông"][2] == "Dịch vụ" and items["Cước dịch vụ viễn thông"][0] == "SP_0000"
    assert items["Văn phòng phẩm"][2] == "Vật tư hàng hóa"
    codes = [r[0] for r in items.values()]
    assert len(codes) == len(set(codes))

    skipped = list(wb["BoQua"].iter_rows(min_row=2, values_only=True))
    assert len(skipped) == 1 and "hủy" in skipped[0][2]
    assert (Path(res.folder) / "HUONG_DAN_NHAP_MISA.txt").is_file()


def test_same_seller_code_for_different_items_gets_unique_codes():
    from hddt.misa import ItemCatalog

    cat = ItemCatalog()
    s = MisaSettings()
    a = cat.code_for("SP.0000", "Xăng RON 95-III", "Lít", "10%", s)
    b = cat.code_for("SP.0000", "Dầu DO 0,05S", "Lít", "10%", s)
    c = cat.code_for("SP.0000", "Xăng RON 95-III", "lít", "10%", s)
    assert a == "SP_0000" and b != a and c == a


def test_export_misa_cli(tmp_path, capsys):
    from hddt import cli

    out = _setup(tmp_path)
    assert cli.main(["--env", "x", "misa", "--thu-muc", str(out)]) == 0
    assert "3 chứng từ" in capsys.readouterr().out


def test_export_misa_empty_dir(tmp_path):
    from hddt import cli

    assert cli.main(["--env", "x", "misa", "--thu-muc", str(tmp_path)]) == 1
