import json
from pathlib import Path

from openpyxl import load_workbook

from hddt.misa import MisaSettings
from hddt.misa_ban import customer_code, export_misa_all, export_misa_ban

FIX = Path(__file__).parent / "fixtures" / "hoadon_mau.xml"
NB = "0312345678"


def _xml(so: str, **repl: str) -> str:
    x = FIX.read_text(encoding="utf-8").replace("<SHDon>125</SHDon>", f"<SHDon>{so}</SHDon>")
    for a, b in repl.items():
        x = x.replace(a, b)
    return x


def _setup(tmp_path: Path) -> Path:
    out = tmp_path / "out"
    sold = out / "xml" / "sold"
    sold.mkdir(parents=True)
    items = []

    def add(so, xml, tthai=1, ttxly=5):
        items.append({"nbmst": NB, "khmshdon": "1", "khhdon": "C25TAA", "shdon": so, "tthai": tthai, "ttxly": ttxly})
        (sold / f"sold_query_{NB}_1_C25TAA_{so}.xml").write_text(xml, encoding="utf-8")

    add("12", _xml("12"))                                     # công ty có MST, 10% + 8%
    add("13", _xml("13", **{"<Ten>CÔNG TY CP XYZ</Ten>": "<Ten></Ten><HVTNMHang>Nguyễn Ngọc Bảo Hân</HVTNMHang>",
                            "<MST>0109876543</MST>": "", "45 Trần Phú, Hà Nội": "12 Lê Lợi, Huế"}))  # cá nhân
    add("14", _xml("14", **{"<TChat>1</TChat><STT>2</STT>": "<TChat>3</TChat><STT>2</STT>"}))  # có chiết khấu
    add("15", _xml("15"), tthai=6)                            # bị hủy
    add("16", _xml("16"), ttxly=2)                            # CQT từ chối
    (out / "danh_sach_banra_20251201_20251231.json").write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")
    return out


def _rows(path):
    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    return header, [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


def test_customer_code():
    assert customer_code("Nguyễn Ngọc Bảo Hân") == "NGUYENNGOCBAOHAN"
    assert customer_code("Đỗ Đức Anh") == "DODUCANH"


def test_export_ban_fills_misa_template(tmp_path):
    out = _setup(tmp_path)
    res = export_misa_ban(out, MisaSettings())
    assert res.vouchers == 3 and res.skipped == 2 and res.customers == 2
    assert Path(res.excel_path).name == "Ban_hang_VND.xlsx"

    wb = load_workbook(res.excel_path)
    assert wb.sheetnames == ["Chứng từ bán hàng"]
    header, rows = _rows(res.excel_path)
    assert header[:3] == ["Hiển thị trên sổ", "Hình thức bán hàng", "Phương thức thanh toán"]
    assert len(rows) == 6  # 3 hóa đơn x 2 dòng hàng (bỏ dòng ghi chú TChat=4)
    for r in rows:
        assert r["Hình thức bán hàng"] == 0 and r["Phương thức thanh toán"] == 1 and r["Lập kèm hóa đơn"] == 1
        assert r["TK Tiền/Chi phí/Nợ (*)"] == "1111" and r["TK Doanh thu/Có (*)"] == "5113"
        assert r["TK thuế GTGT"] == "33311"
    by_so = {}
    for r in rows:
        by_so.setdefault(r["Số hóa đơn"], []).append(r)
    assert by_so["12"][0]["Số chứng từ (*)"] == "BH00012"
    assert by_so["12"][0]["Mã khách hàng"] == "0109876543" and by_so["12"][0]["Mã số thuế"] == "0109876543"
    assert by_so["13"][0]["Mã khách hàng"] == "NGUYENNGOCBAOHAN"
    assert by_so["13"][0]["Tên khách hàng"] == "Nguyễn Ngọc Bảo Hân"
    first = by_so["12"][0]
    assert first["Tên hàng"] == "Dịch vụ tư vấn" and first["ĐVT"] == "Gói"
    assert first["Số lượng"] == 2 and first["Đơn giá"] == 1000000 and first["Thành tiền"] == 2000000
    assert first["% thuế GTGT"] == "10" and first["Tiền thuế GTGT"] == 200000
    assert by_so["12"][1]["% thuế GTGT"] == "8"
    ck = by_so["14"][1]
    assert ck["Thành tiền"] == -500000 and ck["Tiền thuế GTGT"] == -40000
    assert "15" not in by_so and "16" not in by_so

    # Danh mục khách hàng đúng mẫu MISA
    header, kh = _rows(res.dm_kh_path)
    assert "Mã khách hàng (*)" in header
    kh = {r["Mã khách hàng (*)"]: r for r in kh}
    assert kh["0109876543"]["Là tổ chức/cá nhân"] == 0 and kh["0109876543"]["Tên khách hàng (*)"] == "CÔNG TY CP XYZ"
    assert kh["NGUYENNGOCBAOHAN"]["Là tổ chức/cá nhân"] == 1 and kh["NGUYENNGOCBAOHAN"]["Địa chỉ"] == "12 Lê Lợi, Huế"

    # Kiểm tra: 8% bị đánh dấu (mẫu chỉ có 0, 5, 10, KCT); bỏ qua ghi lý do
    wb = load_workbook(res.check_path)
    notes = " ".join(str(r[-1] or "") for r in wb["KiemTra"].iter_rows(min_row=2, values_only=True))
    assert "Thuế suất 8" in notes
    ly_do = [r[2] for r in wb["BoQua"].iter_rows(min_row=2, values_only=True)]
    assert any("hủy" in x for x in ly_do) and any("từ chối" in x for x in ly_do)
    guide = (out / "misa" / "HUONG_DAN_NHAP_MISA.txt").read_text(encoding="utf-8")
    assert "MUA VÀO" in guide and "BÁN RA" in guide


def test_payment_mode_setting(tmp_path):
    out = _setup(tmp_path)
    res = export_misa_ban(out, MisaSettings(ban_thanh_toan="0", tk_doanh_thu="5111"))
    _, rows = _rows(res.excel_path)
    assert {r["Phương thức thanh toán"] for r in rows} == {0}
    assert {r["TK Tiền/Chi phí/Nợ (*)"] for r in rows} == {"131"} and {r["TK Doanh thu/Có (*)"] for r in rows} == {"5111"}


def test_item_code_mapping_is_reused(tmp_path):
    out = _setup(tmp_path)
    export_misa_ban(out, MisaSettings())
    map_path = out / "misa" / "MaHang_MISA.xlsx"
    wb = load_workbook(map_path)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "Bán ra" and row[1].value == "Dịch vụ tư vấn":
            row[3].value = "DV"
        elif row[0].value == "Bán ra" and row[1].value == "Văn phòng phẩm":
            row[3].value = "DV1"
    wb.save(map_path)

    res = export_misa_ban(out, MisaSettings())
    _, rows = _rows(res.excel_path)
    codes = {r["Tên hàng"]: r["Mã hàng (*)"] for r in rows}
    assert codes == {"Dịch vụ tư vấn": "DV", "Văn phòng phẩm": "DV1"}
    # Lần chạy thứ hai không ghi đè mã người dùng đã sửa
    ws = load_workbook(map_path).active
    assert sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "Bán ra") == 2


def test_export_all_only_existing_directions(tmp_path):
    out = _setup(tmp_path)
    allr = export_misa_all(out, MisaSettings())
    assert allr.mua is None and allr.ban and allr.ban.vouchers == 3 and not allr.errors


def test_cli_misa_sold(tmp_path, capsys):
    from hddt import cli

    out = _setup(tmp_path)
    assert cli.main(["--env", "x", "misa", "--thu-muc", str(out)]) == 0
    text = capsys.readouterr().out
    assert "BÁN RA: 3 chứng từ" in text and "MUA VÀO" not in text


def test_combined_list_file_from_pulling_both_directions(tmp_path):
    out = _setup(tmp_path)
    f = out / "danh_sach_banra_20251201_20251231.json"
    f.rename(out / "danh_sach_muavao_banra_20251201_20251231.json")
    res = export_misa_ban(out, MisaSettings())
    assert res.vouchers == 3 and res.skipped == 2
