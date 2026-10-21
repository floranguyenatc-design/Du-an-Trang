"""Tách tiền chưa thuế / tiền thuế theo thuế suất và sheet tổng hợp."""

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from hddt.client import InvoiceRef
from hddt.detail import parsed_from_detail
from hddt.export import INVOICE_COLUMNS, invoice_row, write_workbook
from hddt.thuesuat import SRC_INVOICE, SRC_LINES, SRC_LIST, breakdown, bucket_of
from hddt.xmlparse import InvoiceLine, ParsedInvoice, parse_invoice_file

FIX = Path(__file__).parent / "fixtures" / "hoadon_mau.xml"


def test_bucket_of():
    assert bucket_of("10%") == "10%" and bucket_of("8%") == "8%" and bucket_of("5") == "5%" and bucket_of("0%") == "0%"
    assert bucket_of(0.1) == "10%" and bucket_of("0.08") == "8%"
    assert bucket_of("KCT") == "KCT" and bucket_of("kkknt") == "KKKNT"
    assert bucket_of("KHAC:3.5%") == "Khác" and bucket_of("3,5%") == "Khác"
    assert bucket_of("") is None and bucket_of(None) is None and bucket_of("abc") is None


def test_breakdown_from_invoice_summary_table():
    inv = parse_invoice_file(str(FIX))  # THTTLTSuat: 10% 2.000.000/200.000 ; 8% 500.000/40.000
    sums, src = breakdown({}, inv)
    assert src == SRC_INVOICE
    assert sums == {"10%": [2000000, 200000], "8%": [500000, 40000]}


def test_breakdown_from_list_then_lines():
    raw = {"thttltsuat": [{"tsuat": "10%", "thtien": 1000, "tthue": 100}, {"tsuat": "KCT", "thtien": 500, "tthue": None}]}
    sums, src = breakdown(raw, None)
    assert src == SRC_LIST and sums == {"10%": [1000, 100], "KCT": [500, 0]}

    inv = ParsedInvoice(lines=[
        InvoiceLine(tchat="1", thanh_tien=1000, tien_thue=80, thue_suat="8%"),
        InvoiceLine(tchat="1", thanh_tien=300, tien_thue=30, thue_suat="10%"),
        InvoiceLine(tchat="3", thanh_tien=100, tien_thue=8, thue_suat="8%"),  # chiết khấu: ghi giảm
        InvoiceLine(tchat="4", thanh_tien=None, ten_hang="ghi chú"),
        InvoiceLine(tchat="1", thanh_tien=50, tien_thue=0, thue_suat=""),  # không ghi thuế suất
    ])
    sums, src = breakdown({}, inv)
    assert src == SRC_LINES
    assert sums == {"8%": [900, 72], "10%": [300, 30], "Chưa rõ": [50, 0]}


def test_breakdown_from_portal_detail_json():
    inv = parsed_from_detail({"tgtcthue": 1500, "thttltsuat": [{"tsuat": 0.1, "thtien": 1000, "tthue": 100}, {"tsuat": "KCT", "thtien": 500}]})
    sums, src = breakdown({}, inv)
    assert src == SRC_INVOICE and sums == {"10%": [1000, 100], "KCT": [500, 0]}


def test_invoice_row_rate_columns_period_and_inclusion():
    inv = parse_invoice_file(str(FIX))
    ref = InvoiceRef("purchase", "query", "0312345678", "1", "C25TAA", "125", raw={"tthai": 1, "tgtcthue": 2500000, "tgtthue": 240000})
    row = invoice_row(ref, inv)
    assert row["cth_10"] == 2000000 and row["thue_10"] == 200000 and row["cth_8"] == 500000 and row["thue_8"] == 40000
    assert row["cth_5"] is None and row["cth_kct"] is None and row["cth_chuaro"] is None and row["cth_lech"] is None
    assert row["ky_quy"] == "Quý 4/2025" and row["ky_thang"] == "Tháng 12/2025" and row["tinh_th"] == "Có"
    huy = invoice_row(InvoiceRef("purchase", "query", "1", "1", "A", "2", raw={"tthai": 6, "tgtcthue": 10, "tdlap": "2025-01-05T00:00:00"}), None)
    assert huy["tinh_th"] == "Không" and huy["cth_chuaro"] == 10 and huy["ky_quy"] == "Quý 1/2025"
    lech = invoice_row(InvoiceRef("purchase", "query", "1", "1", "A", "3", raw={"tgtcthue": 2600000}), inv)
    assert lech["cth_lech"] == 100000  # tổng HĐ khác tổng theo thuế suất -> báo lệch


def test_summary_sheet_formulas(tmp_path):
    def row(chieu, ngay, tthai, **amts):
        r = {"chieu": chieu, "ngay_lap": ngay, "shdon": 1, "tinh_th": "Không" if tthai in (4, 6) else "Có",
             "ky_quy": f"Quý {(ngay.month - 1) // 3 + 1}/{ngay.year}", "ky_thang": f"Tháng {ngay.month:02d}/{ngay.year}"}
        r.update(amts)
        return r

    invoices = [
        row("Mua vào", datetime(2024, 1, 5), 1, cth_10=1000, thue_10=100),
        row("Mua vào", datetime(2024, 2, 5), 1, cth_8=500, thue_8=40, cth_kct=20),
        row("Mua vào", datetime(2024, 4, 5), 6, cth_10=999, thue_10=99),  # bị hủy: không tính
        row("Bán ra", datetime(2024, 1, 30), 1, cth_kct=14070000),
    ]
    out = tmp_path / "x.xlsx"
    write_workbook(str(out), invoices, [], [])
    wb = load_workbook(out)
    assert wb.sheetnames == ["HoaDon", "TongHopThueSuat", "ChiTiet", "Loi"]
    ws = wb["TongHopThueSuat"]
    cells = {(ws.cell(r, 1).value, ws.cell(r, 2).value): r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value in ("Mua vào", "Bán ra")}
    # Mua vào có Quý 1 và Quý 2 (Quý 2 chỉ có HĐ hủy nên dòng vẫn có nhưng công thức loại ra), Bán ra có Quý 1
    assert ("Mua vào", "Quý 1/2024") in cells and ("Mua vào", "Quý 2/2024") in cells and ("Bán ra", "Quý 1/2024") in cells
    assert ("Mua vào", "Tháng 02/2024") in cells and ("Mua vào", "Cộng") in cells
    hdr = {ws.cell(6, c).value: c for c in range(1, ws.max_column + 1)}
    col = {key: i for i, (_, key) in enumerate(INVOICE_COLUMNS, start=1)}
    from openpyxl.utils import get_column_letter as L

    f = ws.cell(cells[("Mua vào", "Quý 1/2024")], hdr["Chưa thuế 10%"]).value
    assert f.startswith("=SUMIFS(HoaDon!$" + L(col["cth_10"]) + "$2:")
    assert f"HoaDon!${L(col['chieu'])}$2" in f and f"HoaDon!${L(col['ky_quy'])}$2" in f and f'HoaDon!${L(col["tinh_th"])}$2:${L(col["tinh_th"])}$5,"Có"' in f
    assert ws.cell(cells[("Mua vào", "Quý 1/2024")], hdr["Số hóa đơn"]).value.startswith("=COUNTIFS(")
    assert ws.cell(cells[("Mua vào", "Cộng")], hdr["Cộng tiền chưa thuế"]).value.startswith("=SUM(")
    assert ws["E3"].value.startswith("=COUNTIFS(") and '"Không"' in ws["E3"].value
